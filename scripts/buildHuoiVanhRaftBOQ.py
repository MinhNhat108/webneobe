"""
Bảng thống kê vật tư BÈ PIN hồ Huổi Vanh, mặt bằng 4 bè (bản vẽ 28/10/2026), cùng khuôn với bảng Bản Mòng
(scripts/buildBanMongRaftBOQ.py), kèm bản vẽ sơ đồ hộp 40x60 gia cường tại các vị trí neo:

  python scripts/buildHuoiVanhRaftBOQ.py <bản vẽ bè pin .pdf> <bảng vật tư neo .xlsx> <mẫu hình .xlsx> <out.xlsx> <out.dxf>

- Số tấm pin, phao, inverter, dầm C52, hộp 40x60: lấy NGUYÊN bảng thống kê trên các tờ CT-BE01…04 (bố trí bè)
  và KC-BE-01…04 (khung kết cấu) của "HOHUOIVANH.BỐ TRÍ BÈ PIN 281026.pdf" (script kiểm tra từng số có trên
  đúng tờ đó); phao đỡ cáp ở tờ MB-BP-01.
- Bản vẽ không thể hiện từng thanh thép đủ để đo, nên số cây cấp hàng, mối nối, bulông là ƯỚC LƯỢNG bằng
  công thức Excel theo hệ số ở sheet GiaDinh (cùng hệ số với Bản Mòng, đo trên bè 5 Bản Mòng — hệ MRac G5 như nhau).
- Bản vẽ CHƯA có hộp 40x60 gia cường tai neo: script bố trí cho cả 4 bè (cáp 2 nhánh chữ Y, 2 bích mỗi dây cách
  nhau 3,5 m), vị trí dây theo mặt bằng neo của web, lực dây T của bè theo sheet KiemTra_Nhanh của bảng vật tư
  neo (engine), và vẽ sơ đồ ra <out.dxf> (tọa độ mặt bằng neo, mét).
- Đường bao từng bè đo trên ảnh tờ CT (sheet DoiChieu_MatBang) để so với mặt bằng mà web dùng tính neo.
Hình minh họa ở cột C chép từ <mẫu hình .xlsx> (bảng vật tư bè pin Huổi Vanh 9 bè) theo tên hạng mục.
"""
import io, json, math, sys
from pathlib import Path

import cv2
import ezdxf
import numpy as np
import openpyxl
import pymupdf
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from xlsx_apply_flags import add_apply_flags

ROOT = Path(__file__).resolve().parent.parent
PDF, MOORING_XLSX, TEMPLATE, OUT, OUT_DXF = sys.argv[1:6]
RAFTS = [f'BÈ {i}' for i in range(1, 5)]
SHORT = {n: f'B{n[-1]}' for n in RAFTS}

# ------------------------------------------------------------------ drawing tables
# page (CT sheet), scale, panels, rows of panels, inverters SG-510HX, support floats (count, rows),
# walkway float lines [(count, rows)…], total floats; frame sheet page, C52 (bars, m), box 40x60 (bars, m)
SHEETS = {
    'BÈ 1': dict(page=1, scale=900, pv=2800, dpv=100, inv=4, fs=(2912, 104), fw=[(202, 2), (303, 3)], total=3417, kc=6, c52=(234, '11925.06'), box=(10, '2916.22')),
    'BÈ 2': dict(page=2, scale=1100, pv=9114, dpv=375, inv=13, fs=(9260, 382), fw=[(266, 2), (1538, 7)], total=11064, kc=7, c52=(309, '38745.81'), box=(29, '8718.94')),
    'BÈ 3': dict(page=3, scale=500, pv=1400, dpv=55, inv=2, fs=(1442, 57), fw=[(41, 1), (362, 2)], total=1845, kc=8, c52=(90, '6109.92'), box=(23, '3402.96')),
    'BÈ 4': dict(page=4, scale=1000, pv=5600, dpv=278, inv=8, fs=(5642, 281), fw=[(502, 2), (1310, 3)], total=7454, kc=5, c52=(92, '24310.61'), box=(60, '5161.46')),
}
CABLE_FLOATS = (1209, 2418)  # tờ MB-BP-01, bảng "PHAO ĐỠ CÁP": 1209 hàng, 2418 phao (cả hồ)
# C52 direction on the sheet, from the mean bar length of the table against the raft outline (BÈ 1: 234 bars of
# 51 m across the 50 m width; BÈ 2: 309 × 125 m across the 148 m depth; BÈ 3: 90 × 68 m, only the 90 m and 20 m
# heights fit; BÈ 4: 92 × 264 m along the 348 m length). The box 40x60 runs across the C52.
C52_HORIZONTAL = {'BÈ 1': False, 'BÈ 2': False, 'BÈ 3': False, 'BÈ 4': True}

doc = pymupdf.open(PDF)
for name, s in SHEETS.items():
    text = doc[s['page'] - 1].get_text()
    must = [str(s['pv']), str(s['dpv']), str(s['inv']), str(s['fs'][0]), str(s['fs'][1]), str(s['total']), f"1 : {s['scale']}"]
    must += [str(c) for c, _ in s['fw']]
    missing = [m for m in must if m not in text]
    assert not missing, f'{name}: {missing} không có trên tờ {s["page"]}'
    kc = doc[s['kc'] - 1].get_text()
    assert f'KHUNG KẾT CẤU BÈ PIN {name[-1]}' in kc, f'{name}: tờ {s["kc"]} không phải khung kết cấu'
    missing = [m for m in [str(s['c52'][0]), s['c52'][1], str(s['box'][0]), s['box'][1]] if m not in kc]
    assert not missing, f'{name}: {missing} không có trên tờ {s["kc"]}'
    assert s['fs'][0] + sum(c for c, _ in s['fw']) == s['total'], name
assert str(CABLE_FLOATS[0]) in doc[8].get_text() and str(CABLE_FLOATS[1]) in doc[8].get_text()

# ------------------------------------------------------------------ raft outline measured on the sheet image
def sheet_outline(page, scale, dpi=150):
    p = doc[page - 1]
    clip = pymupdf.Rect(28, 7, 1042, 835)  # drawing area, left of the title block
    pix = p.get_pixmap(dpi=dpi, clip=clip, colorspace=pymupdf.csGRAY)
    a = np.frombuffer(pix.samples, np.uint8).reshape(pix.h, pix.w)
    mpp = 25.4 / 72 * scale / 1000 / (dpi / 72)
    b = (a < 200).astype(np.uint8)
    b = cv2.morphologyEx(b, cv2.MORPH_CLOSE, np.ones((max(3, int(1.5 / mpp)),) * 2, np.uint8))  # fill the raft
    b = cv2.morphologyEx(b, cv2.MORPH_OPEN, np.ones((max(5, int(6 / mpp)),) * 2, np.uint8))    # drop tables, dimensions
    n, lab, st, _ = cv2.connectedComponentsWithStats(b, 8)
    m = (lab == 1 + int(np.argmax(st[1:, 4]))).astype(np.uint8)
    cnt = max(cv2.findContours(m, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)[0], key=cv2.contourArea)
    x, y, w, h = cv2.boundingRect(cnt)
    return cv2.contourArea(cnt) * mpp * mpp, w * mpp, h * mpp

# ------------------------------------------------------------------ web layout (what the mooring is computed on)
data = ROOT / 'src' / 'data'
polys = {r['name']: r for r in json.load(open(data / 'huoiVanhRaftPolygons_v2.json', encoding='utf-8'))}
catalogue = {r['name']: r for r in json.load(open(data / 'huoiVanhRaftCatalogue.json', encoding='utf-8'))}
lines = json.load(open(data / 'huoiVanhCoordinates_v2.json', encoding='utf-8'))

def principal(points):
    """Dominant edge direction (mod 90°) of a rectilinear raft, and the extents along it and across it."""
    sx = sy = 0.0
    for i in range(len(points)):
        (x1, y1), (x2, y2) = points[i], points[(i + 1) % len(points)]
        L = math.hypot(x2 - x1, y2 - y1); a = 4 * math.atan2(y2 - y1, x2 - x1)
        sx += L * math.cos(a); sy += L * math.sin(a)
    th = math.atan2(sy, sx) / 4
    u = (math.cos(th), math.sin(th)); v = (-u[1], u[0])
    pu = [x * u[0] + y * u[1] for x, y in points]; pv = [x * v[0] + y * v[1] for x, y in points]
    return th, max(pu) - min(pu), max(pv) - min(pv)

layout = {}
for name in RAFTS:
    pts = [(p['x'], p['y']) for p in polys[name]['points']]
    th, Wu, Wv = principal(pts)
    s = SHEETS[name]
    area, w, h = sheet_outline(s['page'], s['scale'])
    sheet_h_is_u = abs(w - Wu) + abs(h - Wv) <= abs(w - Wv) + abs(h - Wu)
    box_along_sheet_h = not C52_HORIZONTAL[name]
    box_dir = th if box_along_sheet_h == sheet_h_is_u else th + math.pi / 2
    ccw = sum(x1 * y2 - x2 * y1 for (x1, y1), (x2, y2) in zip(pts, pts[1:] + pts[:1])) > 0
    layout[name] = dict(pts=pts, th=th, Wu=Wu, Wv=Wv, area=area, w=w, h=h, box_dir=box_dir, sheet_h_is_u=sheet_h_is_u, ccw=ccw)

# ------------------------------------------------------------------ T per raft from the mooring BOQ (engine)
mw = openpyxl.load_workbook(MOORING_XLSX, data_only=True)['KiemTra_Nhanh']
T = {r[0]: r[1] for r in mw.iter_rows(min_row=4, values_only=True) if r[0] in RAFTS}
assert set(T) == set(RAFTS), T

# ------------------------------------------------------------------ anchor-point reinforcement clusters
CLEAT_GAP, L_GC, B_C52, P_KEP, K_Y, GAMMA, BAR_GAP, INSET = 3.5, 15.0, 1.225, 7.0, 0.7, 1.2, 0.6, 1.0

def nearest_edge(pts, x, y):
    best = None
    for i in range(len(pts)):
        (x1, y1), (x2, y2) = pts[i], pts[(i + 1) % len(pts)]
        dx, dy = x2 - x1, y2 - y1; L2 = dx * dx + dy * dy
        t = max(0, min(1, ((x - x1) * dx + (y - y1) * dy) / L2))
        d = math.hypot(x1 + t * dx - x, y1 + t * dy - y)
        if best is None or d < best[0]:
            best = (d, i, t * math.sqrt(L2), math.atan2(dy, dx))
    return best

def bars_needed(cleats, t):
    """Same as the sheet formula with the GiaDinh defaults: MAX(1; ROUNDUP(bích × K_Y × T × γ / (giao C52 × P_KEP)))."""
    crossings = math.floor(L_GC / B_C52) + 1
    return max(1, math.ceil(cleats * K_Y * t * GAMMA / (crossings * P_KEP) - 0.0005))

clusters = []  # (raft, id, kind, codes, geometry)
for name in RAFTS:
    lay = layout[name]; pts = lay['pts']
    inward = lambda ex, ey: (-ey, ex) if lay['ccw'] else (ey, -ex)
    bdir = (math.cos(lay['box_dir']), math.sin(lay['box_dir']))
    along, across = {}, []
    for ln in [l for l in lines if l['raft'] == name]:
        d, edge, pos, ang = nearest_edge(pts, ln['xRaft'], ln['yRaft'])
        diff = abs((ang - lay['box_dir'] + math.pi / 2) % math.pi - math.pi / 2)
        for k, off in (('a', -CLEAT_GAP / 2), ('b', CLEAT_GAP / 2)):
            if diff < math.pi / 4:
                along.setdefault(edge, []).append((pos + off, ln['code'] + k))
            else:
                across.append((edge, pos + off, ln['code'] + k))
    k = 0
    for edge in sorted(along):
        (x1, y1), (x2, y2) = pts[edge], pts[(edge + 1) % len(pts)]
        Le = math.hypot(x2 - x1, y2 - y1); ex, ey = (x2 - x1) / Le, (y2 - y1) / Le; nx, ny = inward(ex, ey)
        groups, group = [], []
        for pos, code in sorted(along[edge]):
            if group and pos - group[0][0] > L_GC - 1.0:
                groups.append(group); group = []
            group.append((pos, code))
        if group: groups.append(group)
        for g in groups:
            k += 1
            c = (g[0][0] + g[-1][0]) / 2
            a, b = c - L_GC / 2, c + L_GC / 2
            if Le >= L_GC + 1:  # keep the bar on its edge
                shift = max(0.5 - a, 0) - max(b - (Le - 0.5), 0); a += shift; b += shift
            else:
                a, b = 0.5, Le - 0.5
            geo = dict(p0=(x1 + ex * a + nx * INSET, y1 + ey * a + ny * INSET), p1=(x1 + ex * b + nx * INSET, y1 + ey * b + ny * INSET),
                       step=(nx * BAR_GAP, ny * BAR_GAP), cleats=[(x1 + ex * p, y1 + ey * p) for p, _ in g])
            clusters.append((name, f'{SHORT[name]}-{k:02d}', 'Dọc cạnh (hộp song song mép bè, cách mép 1 hàng phao)', [c for _, c in g], geo))
    for edge, pos, code in across:
        (x1, y1), (x2, y2) = pts[edge], pts[(edge + 1) % len(pts)]
        Le = math.hypot(x2 - x1, y2 - y1); ex, ey = (x2 - x1) / Le, (y2 - y1) / Le; nx, ny = inward(ex, ey)
        ux, uy = bdir if bdir[0] * nx + bdir[1] * ny > 0 else (-bdir[0], -bdir[1])
        cx, cy = x1 + ex * pos, y1 + ey * pos
        k += 1
        geo = dict(p0=(cx + ux * 0.5, cy + uy * 0.5), p1=(cx + ux * (0.5 + L_GC), cy + uy * (0.5 + L_GC)), step=(ex * BAR_GAP, ey * BAR_GAP), cleats=[(cx, cy)])
        clusters.append((name, f'{SHORT[name]}-{k:02d}', 'Vuông góc cạnh (hộp chạy từ mép vào trong, cùng phương lực dây)', [code], geo))

# ------------------------------------------------------------------ workbook
wb = openpyxl.Workbook()
thin = Side(style='thin'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='F2F2F2'); SEC = PatternFill('solid', fgColor='D9D9D9'); INPUT = PatternFill('solid', fgColor='FFFF99')
F = Font(name='Calibri', size=12); FB = Font(name='Calibri', size=12, bold=True)
WRAP = Alignment(wrap_text=True, vertical='center'); CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
RC = 'CDEF'  # raft columns of DuLieuBanVe
TC = 'EFGH'  # raft columns of TongHop

def grid(ws, r0, r1, c1):
    for row in ws.iter_rows(min_row=r0, max_row=r1, max_col=c1):
        for c in row:
            c.border = BORDER
            if not c.font.b: c.font = F
            if c.alignment.horizontal is None: c.alignment = WRAP

# ---- GiaDinh
gd = wb.active; gd.title = 'GiaDinh'
GD = [
    ('L_C52', 'Chiều dài cây dầm C52 khi cấp hàng', 6, 'm', 'GIẢ ĐỊNH. HD lắp Mibet G5 dùng đoạn L2990 / L3750 — nếu cấp đoạn ~3 m thì số mối nối gần gấp đôi'),
    ('L_H', 'Chiều dài cây hộp 40x60', 6, 'm', 'GIẢ ĐỊNH, cây thép hộp thương mại 6 m'),
    ('M_C52', 'Khối lượng dầm C52', 2.36, 'kg/m', 'GIẢ ĐỊNH tiết diện C 52×41×12×2,0 (chưa có catalogue nhà sản xuất)'),
    ('M_H', 'Khối lượng hộp 40x60', 3.01, 'kg/m', 'GIẢ ĐỊNH 40×60×2,0'),
    ('K_FS', 'Gá trước (Front Support) mỗi tấm pin', 2, 'cái/tấm', 'ƯỚC LƯỢNG theo HD lắp G5'),
    ('K_BS', 'Gá sau (Back Support) mỗi tấm pin', 2, 'cái/tấm', 'ƯỚC LƯỢNG'),
    ('K_EC', 'Bộ kẹp biên tấm pin (End Clamp Kit) mỗi tấm', 2, 'bộ/tấm', 'ƯỚC LƯỢNG'),
    ('K_GC', 'Bộ kẹp tiếp địa mỗi tấm', 2, 'bộ/tấm', 'ƯỚC LƯỢNG; hỏi Mibet: kẹp biên/kẹp giữa có răng tiếp địa thì bỏ được hạng mục này, hoặc 1 bộ/tấm'),
    ('K_RN8', 'Bộ đai ốc thoi M8×25 bắt gá vào C52', 4, 'bộ/tấm', 'ƯỚC LƯỢNG = 1 bộ mỗi gá'),
    ('K_FB', 'Bộ bulông nhựa nối phao', 2, 'bộ/phao', 'ƯỚC LƯỢNG (HD: các phao nối với nhau bằng bulông + đai ốc nhựa)'),
    ('HH', 'Hao hụt phụ kiện, bulông', 3, '%', 'GIẢ ĐỊNH, chỉ cộng ở cột "Tổng có hao hụt"'),
    ('L_GC', 'Chiều dài 1 thanh hộp 40x60 gia cường tai neo', L_GC, 'm', 'Theo bản vẽ gia cường bè 1 Huổi Vanh 9 bè (13,4–15,5 m)'),
    ('P_KEP', 'Sức chịu 1 móc "几" (2 đai ốc thoi M10×30 trong rãnh C52)', P_KEP, 'kN', 'GIẢ ĐỊNH — chưa có số liệu nhà sản xuất; quyết định số thanh xếp song song ở cụm chịu lực lớn. Bản vẽ GC-01 vẽ số thanh theo giá trị mặc định này'),
    ('GAMMA', 'Hệ số tải trọng cho lực dây', GAMMA, '', 'Như kiểm tra cọc (pileBendingLoadFactor)'),
    ('K_Y', 'Phần lực dây vào MỘT bích (cáp 2 nhánh chữ Y)', K_Y, '× T', 'Chia đều ở góc 60°: 0,58 T; lấy 0,7 T vì 2 nhánh có thể dài lệch nhau (cùng giả định với bảng vật tư neo, sheet KiemTra_Nhanh)'),
    ('K_DU', 'Phần thừa trung bình của cây cuối mỗi thanh', 0.5, 'cây/thanh', 'ƯỚC LƯỢNG: không đo được từng thanh → số cây = ROUNDUP(tổng dài / L + số thanh × K_DU)'),
    ('K_XF', 'Số chỗ dầm C52 vượt qua một hàng phao, trên 1 m dầm', 0.522, 'điểm/m', 'ĐO trên bè 5 Bản Mòng (cùng hệ G5). Kiểm lại trên BÈ 1 Huổi Vanh: 234 dầm × 26 hàng phao mỗi khối / 11.925 m = 0,51'),
    ('B_C52', 'Khoảng cách trung bình giữa các dầm C52 (theo phương hộp)', B_C52, 'm', 'ĐO trên bè 5 Bản Mòng; Huổi Vanh: bề rộng / số dầm = 1,24 m (BÈ 1, BÈ 2), 1,22 m (BÈ 3). 1 bulông M10×120 mỗi giao điểm hộp – dầm'),
]
gd['A1'] = 'THÔNG SỐ GIẢ ĐỊNH — ô tô vàng là ô NHẬP, sửa là cả file tự tính lại (bản vẽ GC-01 vẽ theo giá trị mặc định)'; gd['A1'].font = Font(name='Calibri', size=13, bold=True)
for j, h in enumerate(['Mã', 'Thông số', 'Giá trị', 'Đơn vị', 'Ghi chú']):
    c = gd.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
G = {}
for i, (code, label, val, unit, note) in enumerate(GD):
    r = 4 + i; G[code] = f'GiaDinh!$C${r}'
    for j, v in enumerate([code, label, val, unit, note]):
        gd.cell(r, j + 1, v)
    gd.cell(r, 3).fill = INPUT
grid(gd, 3, 3 + len(GD), 5)
for col, w in zip('ABCDE', [9, 52, 10, 10, 110]): gd.column_dimensions[col].width = w

# ---- DuLieuBanVe
dl = wb.create_sheet('DuLieuBanVe', 0)
dl['A1'] = 'SỐ LIỆU LẤY TỪ BẢN VẼ "HOHUOIVANH.BỐ TRÍ BÈ PIN 281026" (bảng thống kê trên các tờ CT-BE01…04, KC-BE-01…04, MB-BP-01)'; dl['A1'].font = Font(name='Calibri', size=13, bold=True)
for j, h in enumerate(['Mã', 'Hạng mục'] + RAFTS + ['TỔNG', 'Nguồn']):
    c = dl.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
fw = lambda name, k: SHEETS[name]['fw'][k][0] if k < len(SHEETS[name]['fw']) else 0
lineCount = {n: catalogue[n]['cableCount'] for n in RAFTS}
DL = [
    ('PV', 'Tấm pin (tấm)', lambda n: SHEETS[n]['pv'], 'Bảng "SỐ TẤM PIN MẶT TRỜI THEO BÈ" (tờ CT-BE), cột Số tấm pin'),
    ('DPV', 'Dải pin (hàng)', lambda n: SHEETS[n]['dpv'], 'Như trên, cột SỐ LƯỢNG HÀNG'),
    ('I510', 'Inverter SG-510HX (bộ)', lambda n: SHEETS[n]['inv'], 'Bảng "INVESTER" (tờ CT-BE)'),
    ('FS', 'Phao đỡ pin — Main Floater (cái)', lambda n: SHEETS[n]['fs'][0], 'Bảng "SỐ PHAO NỔI THEO BÈ", dòng "DẢI PHAO ĐỠ PIN"'),
    ('FSR', 'Số hàng phao đỡ pin', lambda n: SHEETS[n]['fs'][1], 'Như trên, cột SỐ LƯỢNG HÀNG'),
    ('FWA', 'Phao đi lại nhóm 1 (cái)', lambda n: fw(n, 0), 'Dòng "DẢI PHAO ĐI LẠI" thứ nhất của bảng phao'),
    ('FWB', 'Phao đi lại nhóm 2 (cái)', lambda n: fw(n, 1), 'Dòng "DẢI PHAO ĐI LẠI" thứ hai'),
    ('FWR', 'Số hàng phao đi lại', lambda n: sum(r for _, r in SHEETS[n]['fw']), 'Cộng cột SỐ LƯỢNG HÀNG các dòng phao đi lại'),
    ('FT', 'Tổng phao (cái)', None, 'Khớp "Grand total" của bảng phao từng tờ'),
    ('C52N', 'Dầm C52 — số thanh', lambda n: SHEETS[n]['c52'][0], 'Bảng "THỐNG KÊ KHUNG THÉP" (tờ KC-BE)'),
    ('C52L', 'Dầm C52 — chiều dài (m)', lambda n: float(SHEETS[n]['c52'][1]), 'Bảng "THỐNG KÊ KHUNG THÉP" (tờ KC-BE)'),
    ('HN', 'Hộp 40x60 kết cấu — số thanh', lambda n: SHEETS[n]['box'][0], 'Bảng "THỐNG KÊ KHUNG THÉP" (tờ KC-BE)'),
    ('HL', 'Hộp 40x60 kết cấu — chiều dài (m)', lambda n: float(SHEETS[n]['box'][1]), 'Bảng "THỐNG KÊ KHUNG THÉP" (tờ KC-BE). BÈ 3: trung bình 148 m/thanh, dài hơn kích thước lớn nhất của bè (109,9 m) — cần kiểm lại bảng tờ KC-BE-03'),
    ('LINE', 'Số dây neo của bè (tuyến cáp)', lambda n: lineCount[n], 'Mặt bằng neo trên web (bảng vật tư neo Huổi Vanh)'),
]
D = {}
for i, (code, label, fn, src) in enumerate(DL):
    r = 4 + i; D[code] = r
    dl.cell(r, 1, code); dl.cell(r, 2, label); dl.cell(r, 8, src)
    for j, n in enumerate(RAFTS):
        col = RC[j]
        dl[f'{col}{r}'] = f'={col}{D["FS"]}+{col}{D["FWA"]}+{col}{D["FWB"]}' if code == 'FT' else fn(n)
        dl[f'{col}{r}'].number_format = '#,##0.00' if code in ('C52L', 'HL') else '#,##0'
    dl[f'G{r}'] = f'=SUM(C{r}:F{r})'; dl[f'G{r}'].number_format = dl[f'C{r}'].number_format
r = 4 + len(DL); D['FC'] = r
dl.cell(r, 1, 'FC'); dl.cell(r, 2, 'Phao đỡ cáp điện giữa các bè (cái, cả hồ)'); dl[f'G{r}'] = CABLE_FLOATS[1]
dl.cell(r, 8, f'Tờ MB-BP-01, bảng "PHAO ĐỠ CÁP": {CABLE_FLOATS[0]} hàng, {CABLE_FLOATS[1]} phao (không chia theo bè)')
grid(dl, 3, r, 8)
for col, w in zip('ABCDEFGH', [8, 44, 11, 11, 11, 11, 12, 80]): dl.column_dimensions[col].width = w

# ---- GiaCuong_TaiNeo
gc = wb.create_sheet('GiaCuong_TaiNeo')
gc['A1'] = 'HỘP 40x60 GIA CƯỜNG TAI NEO — BỐ TRÍ CHO CẢ 4 BÈ (bản vẽ 28/10 chưa có), CÁP 2 NHÁNH CHỮ Y (2 bích mỗi dây, cách nhau 3,5 m). Sơ đồ vị trí: bản vẽ GC-01'
gc['A2'] = ('Mỗi cụm tai neo một (hoặc vài) thanh hộp dài L_GC kẹp móc "几" vào mọi dầm C52 nó cắt qua. Số thanh = MAX(1; số bích × K_Y × T × γ / (số giao C52 × P_KEP)). '
            'Bích trên cạnh song song với hộp 40x60 → các bích trong đoạn L_GC − 1 m dùng chung 1 thanh dọc cạnh; bích trên cạnh vuông góc → mỗi bích 1 thanh từ mép vào trong. '
            'Vị trí dây theo mặt bằng neo của web; hướng hộp 40x60 của từng bè theo tờ KC-BE (28/10). T là lực dây của bè (như cáp, ma-nê, tấm chia Y trong bảng vật tư neo).')
for a in ('A1', 'A2'): gc[a].font = FB if a == 'A1' else F
gc.merge_cells('A2:S2'); gc['A2'].alignment = WRAP; gc.row_dimensions[2].height = 48
HDR = ['Bè', 'Mã cụm (bản vẽ GC-01)', 'Cách bố trí', 'Bích trong cụm (mã dây + a/b)', 'Số bích', 'T dây (kN, gió 20 m/s)', 'Lực thiết kế cụm = số bích × K_Y × T × γ (kN)', 'Khoảng cách dầm C52 (m)', 'Dài 1 thanh (m)',
       'Giao C52 mỗi thanh', 'Sức chịu 1 thanh (kN)', 'Số thanh', 'Tổng dài (m)', 'Số cây cấp hàng', 'Số mối nối', 'Móc "几" (cái)', 'Đai ốc thoi M10×30 (bộ)', 'Bulông M10×120 (bộ)', 'Khối lượng (kg)']
for j, h in enumerate(HDR):
    c = gc.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
gc.row_dimensions[3].height = 62
for i, (name, cid, kind, codes, _) in enumerate(clusters):
    r = 4 + i
    vals = [name, cid, kind, ', '.join(codes), len(codes), T[name],
            f'=E{r}*F{r}*{G["K_Y"]}*{G["GAMMA"]}', f'={G["B_C52"]}', f'={G["L_GC"]}', f'=ROUNDDOWN(I{r}/H{r},0)+1', f'=J{r}*{G["P_KEP"]}',
            f'=MAX(1,ROUNDUP(G{r}/K{r}-0.0005,0))', f'=L{r}*I{r}', f'=L{r}*ROUNDUP(I{r}/{G["L_H"]}-0.0005,0)', f'=L{r}*(ROUNDUP(I{r}/{G["L_H"]}-0.0005,0)-1)',
            f'=L{r}*J{r}', f'=2*P{r}', f'=P{r}', f'=M{r}*{G["M_H"]}']
    for j, v in enumerate(vals): gc.cell(r, j + 1, v)
    gc[f'G{r}'].number_format = '0.0'; gc[f'K{r}'].number_format = '0.0'; gc[f'M{r}'].number_format = '0.0'; gc[f'S{r}'].number_format = '#,##0.0'
GC_LAST = 3 + len(clusters)
grid(gc, 3, GC_LAST, 19)
for col, w in zip('ABCDEFGHIJKLMNOPQRS', [8, 10, 30, 34, 7, 10, 14, 10, 8, 9, 10, 8, 9, 9, 8, 9, 10, 10, 11]): gc.column_dimensions[col].width = w

# ---- DoiChieu_MatBang
dc = wb.create_sheet('DoiChieu_MatBang')
dc['A1'] = 'ĐỐI CHIẾU TỪNG TỜ CT-BE (28/10/2026) VỚI MẶT BẰNG BÈ MÀ WEB DÙNG TÍNH NEO (bản vẽ CAD 27/10/2026)'; dc['A1'].font = Font(name='Calibri', size=13, bold=True)
dc['A2'] = ('Đường bao bè đo trên ảnh từng tờ CT-BE (tỷ lệ in trên tờ, gồm cả mép phao). Số tấm pin trên web lấy đúng bảng từng tờ (18.914 tấm, từ 28/10/2026; '
            'trước đó tạm 18.354 tấm của mặt bằng 08/10 chia theo diện tích).')
dc.merge_cells('A2:N2'); dc['A2'].alignment = WRAP; dc.row_dimensions[2].height = 34
H2 = ['Bè', 'Tỷ lệ tờ', 'Diện tích bản vẽ 28/10 (m²)', 'Diện tích trên web (m²)', 'Lệch diện tích', 'Kích thước bản vẽ (m)', 'Kích thước trên web (m)', 'Tấm pin bản vẽ', 'Tấm pin trên web', 'Lệch tấm pin',
      'Dầm C52 TB 1 thanh (m)', 'Hộp 40x60 TB 1 thanh (m)', 'Phương dầm C52 trên tờ', 'Số dây neo']
for j, h in enumerate(H2):
    c = dc.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
dc.row_dimensions[3].height = 48
for i, n in enumerate(RAFTS):
    r = 4 + i; lay = layout[n]; s = SHEETS[n]
    dims_w = (lay['Wu'], lay['Wv']) if lay['sheet_h_is_u'] else (lay['Wv'], lay['Wu'])
    row = [n, f"1:{s['scale']}", round(lay['area']), polys[n]['area_m2'], f'=C{r}/D{r}-1', f"{lay['w']:.1f} × {lay['h']:.1f}".replace('.', ','), f"{dims_w[0]:.1f} × {dims_w[1]:.1f}".replace('.', ','),
           s['pv'], catalogue[n]['solarPanelCount'], f'=H{r}/I{r}-1', f"=DuLieuBanVe!{RC[i]}{D['C52L']}/DuLieuBanVe!{RC[i]}{D['C52N']}", f"=DuLieuBanVe!{RC[i]}{D['HL']}/DuLieuBanVe!{RC[i]}{D['HN']}",
           'ngang' if C52_HORIZONTAL[n] else 'đứng', lineCount[n]]
    for j, v in enumerate(row): dc.cell(r, j + 1, v)
    for col in 'EJ': dc[f'{col}{r}'].number_format = '+0.0%;-0.0%'
    for col in 'KL': dc[f'{col}{r}'].number_format = '0.0'
    for col in 'CDHI': dc[f'{col}{r}'].number_format = '#,##0'
r = 8
dc.cell(r, 1, 'TỔNG'); dc[f'C{r}'] = '=SUM(C4:C7)'; dc[f'D{r}'] = '=SUM(D4:D7)'; dc[f'E{r}'] = f'=C{r}/D{r}-1'; dc[f'H{r}'] = '=SUM(H4:H7)'; dc[f'I{r}'] = '=SUM(I4:I7)'; dc[f'J{r}'] = f'=H{r}/I{r}-1'; dc[f'N{r}'] = '=SUM(N4:N7)'
for col in 'EJ': dc[f'{col}{r}'].number_format = '+0.0%;-0.0%'
for col in 'CDHI': dc[f'{col}{r}'].number_format = '#,##0'
grid(dc, 3, 8, 14)
for c in dc[8]: c.font = FB
for col, w in zip('ABCDEFGHIJKLMN', [8, 8, 12, 12, 10, 16, 16, 10, 11, 9, 11, 11, 10, 9]): dc.column_dimensions[col].width = w

# ---- TongHop
th = wb.create_sheet('TongHop', 0)
th['A1'] = 'BẢNG TỔNG HỢP VẬT TƯ BÈ PIN — 4 BÈ HỒ HUỔI VANH (bản vẽ "HOHUOIVANH.BỐ TRÍ BÈ PIN 281026", HD lắp Mibet MRac G5)'
th['A2'] = ('Nguồn: BV = lấy đúng bảng thống kê trên bản vẽ; ƯL = ước lượng theo thông số sheet GiaDinh (phải đối chiếu BOM nhà sản xuất); '
            'NEO = theo mặt bằng neo trên web. Bản vẽ CHƯA có hộp 40x60 gia cường tai neo: mục D bố trí theo sơ đồ GC-01 (ước tính).')
th.merge_cells('A1:L1'); th.merge_cells('A2:L2')
th['A1'].font = Font(name='Calibri', size=13, bold=True); th['A2'].alignment = WRAP; th.row_dimensions[1].height = 17; th.row_dimensions[2].height = 34
for j, h in enumerate(['STT', 'Hạng mục', 'Hình minh họa (tài liệu Mibet)', 'Đơn vị'] + RAFTS + ['TỔNG', 'Tổng có hao hụt', 'Nguồn', 'Ghi chú']):
    c = th.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
th.row_dimensions[3].height = 31

def dlref(code, j): return f'DuLieuBanVe!{RC[j]}{D[code]}'
def gcsum(col, n): return f'SUMIF(GiaCuong_TaiNeo!$A$4:$A${GC_LAST},"{n}",GiaCuong_TaiNeo!${col}$4:${col}${GC_LAST})'
ITEMS = [
    ('A. PIN, INVERTER, PHAO NỔI',),
    ('Tấm pin mặt trời', 'tấm', lambda j: '=' + dlref('PV', j), 'same', 'BV', '18.914 tấm theo bảng từng tờ CT-BE (28/10/2026); lực gió và lực dây trên web tính theo số tấm từng bè này'),
    ('Inverter SG-510HX', 'bộ', lambda j: '=' + dlref('I510', j), 'same', 'BV', '27 inverter; bản vẽ không có SG350HX'),
    ('Phao đỡ pin (Main Floater G5G)', 'cái', lambda j: '=' + dlref('FS', j), 'same', 'BV', ''),
    ('Phao đi lại (Walkway Floater) — nhóm 1', 'cái', lambda j: '=' + dlref('FWA', j), 'same', 'BV', 'Bản vẽ chia phao đi lại thành 2 dòng; cần nhà sản xuất xác nhận là 1 hay 2 loại phao'),
    ('Phao đi lại (Walkway Floater) — nhóm 2', 'cái', lambda j: '=' + dlref('FWB', j), 'same', 'BV', ''),
    ('Tổng phao', 'cái', lambda j: '=' + dlref('FT', j), 'same', 'BV', ''),
    ('Phao đỡ cáp điện giữa các bè', 'cái', None, 'cable', 'BV', f'Tờ MB-BP-01: {CABLE_FLOATS[0]} hàng × 2 = {CABLE_FLOATS[1]} phao, cả hồ (không chia theo bè)'),
    ('Bộ bulông nhựa nối phao', 'bộ', lambda j: f'=ROUNDUP({dlref("FT", j)}*{G["K_FB"]},0)', 'hh', 'ƯL', 'Tổng phao bè × hệ số K_FB ở GiaDinh (chưa gồm phao đỡ cáp)'),
    ('B. DẦM C52 (đặt trên rãnh phao)',),
    ('Dầm C52 — số thanh liên tục (bảng bản vẽ)', 'thanh', lambda j: '=' + dlref('C52N', j), 'same', 'BV', ''),
    ('Dầm C52 — chiều dài', 'm', lambda j: '=' + dlref('C52L', j), 'same', 'BV', 'Bảng tờ KC-BE'),
    ('Dầm C52 — số cây cấp hàng', 'cây', lambda j: f'=ROUNDUP({dlref("C52L", j)}/{G["L_C52"]}+{dlref("C52N", j)}*{G["K_DU"]},0)', 'same', 'ƯL', 'Không đo được từng thanh: tổng dài / L_C52 + số thanh × K_DU (GiaDinh)'),
    ('Bộ nối dầm C52 (Splice for Beam-C)', 'bộ', lambda j: f'={TC[j]}{{r_prev}}-{TC[j]}{{r_bars}}', 'hh', 'ƯL', 'Số cây − số thanh. Nối đối đầu; mỗi bộ: 1 bản nối + 2 bulông M10×70 + 2 chốt M10×60 + đai ốc, long đen (tài liệu phụ kiện Mibet)'),
    ('Dầm C52 — khối lượng', 'kg', lambda j: f'={dlref("C52L", j)}*{G["M_C52"]}', 'same', 'ƯL', 'Theo tiết diện GIẢ ĐỊNH'),
    ('Bulông M8×35 bắt dầm C52 vào phao', 'bộ', lambda j: f'=ROUNDUP({dlref("C52L", j)}*{G["K_XF"]},0)', 'hh', 'ƯL', '1 bộ tại mỗi chỗ dầm C52 vượt qua một hàng phao (K_XF)'),
    ('C. HỘP 40x60 KẾT CẤU (nằm ngang trên dầm C52)',),
    ('Hộp 40x60 — số thanh liên tục', 'thanh', lambda j: '=' + dlref('HN', j), 'same', 'BV', ''),
    ('Hộp 40x60 — chiều dài', 'm', lambda j: '=' + dlref('HL', j), 'same', 'BV', 'Bảng tờ KC-BE. BÈ 3: 23 thanh, trung bình 148 m/thanh > 109,9 m (kích thước lớn nhất của bè) — CẦN KIỂM LẠI bảng KC-BE-03'),
    ('Hộp 40x60 — số cây cấp hàng', 'cây', lambda j: f'=ROUNDUP({dlref("HL", j)}/{G["L_H"]}+{dlref("HN", j)}*{G["K_DU"]},0)', 'same', 'ƯL', ''),
    ('Bộ nối hộp 40x60 (Splice for Square Steel)', 'bộ', lambda j: f'={TC[j]}{{r_prev}}-{TC[j]}{{r_bars}}', 'hh', 'ƯL', 'Nối đối đầu bằng bộ nối của nhà sản xuất (datasheet G5G); tài liệu Mibet không dùng nối chồng'),
    ('Hộp 40x60 — khối lượng', 'kg', lambda j: f'={dlref("HL", j)}*{G["M_H"]}', 'same', 'ƯL', 'Theo tiết diện GIẢ ĐỊNH'),
    ('Bulông M10×120 bắt hộp 40x60 vào dầm C52', 'bộ', lambda j: f'=ROUNDUP({dlref("HL", j)}/{G["B_C52"]},0)', 'hh', 'ƯL', '1 bộ tại mỗi giao điểm hộp – dầm C52 (HD lắp G5 bước 8)'),
    ('D. HỘP 40x60 GIA CƯỜNG TAI NEO (bản vẽ bè CHƯA có — bố trí theo sơ đồ GC-01, ƯỚC TÍNH)',),
    ('Hộp 40x60 gia cường — số thanh', 'thanh', lambda j: '=' + gcsum('L', RAFTS[j]), 'same', 'ƯL', 'Sheet GiaCuong_TaiNeo (từng cụm tai neo, mã cụm như bản vẽ GC-01)'),
    ('Hộp 40x60 gia cường — chiều dài', 'm', lambda j: '=' + gcsum('M', RAFTS[j]), 'same', 'ƯL', ''),
    ('Hộp 40x60 gia cường — số cây cấp hàng', 'cây', lambda j: '=' + gcsum('N', RAFTS[j]), 'same', 'ƯL', ''),
    ('Bộ nối hộp 40x60 (gia cường)', 'bộ', lambda j: '=' + gcsum('O', RAFTS[j]), 'hh', 'ƯL', ''),
    ('Móc chữ "几" kẹp hộp vào dầm C52', 'cái', lambda j: '=' + gcsum('P', RAFTS[j]), 'hh', 'ƯL', 'HD lắp G5 bước 8: hộp dùng làm điểm neo kẹp bằng móc "几" tại mọi giao điểm với C52'),
    ('Bộ đai ốc thoi M10×30 (cho móc "几")', 'bộ', lambda j: '=' + gcsum('Q', RAFTS[j]), 'hh', 'ƯL', '2 bộ mỗi móc'),
    ('Bulông M10×120 (hộp gia cường)', 'bộ', lambda j: '=' + gcsum('R', RAFTS[j]), 'hh', 'ƯL', '1 bộ mỗi móc'),
    ('Hộp 40x60 gia cường — khối lượng', 'kg', lambda j: '=' + gcsum('S', RAFTS[j]), 'same', 'ƯL', 'Theo tiết diện GIẢ ĐỊNH'),
    ('E. GÁ ĐỠ VÀ KẸP TẤM PIN',),
    ('Gá trước (Front Support G5)', 'cái', lambda j: f'={dlref("PV", j)}*{G["K_FS"]}', 'hh', 'ƯL', ''),
    ('Gá sau (Back Support G5)', 'cái', lambda j: f'={dlref("PV", j)}*{G["K_BS"]}', 'hh', 'ƯL', ''),
    ('Bộ kẹp biên tấm pin (End Clamp Kit, bulông M8×30)', 'bộ', lambda j: f'={dlref("PV", j)}*{G["K_EC"]}', 'hh', 'ƯL', ''),
    ('Bộ kẹp tiếp địa tấm pin', 'bộ', lambda j: f'={dlref("PV", j)}*{G["K_GC"]}', 'hh', 'ƯL', 'Cần nối đẳng thế khung tấm pin; hỏi Mibet kẹp của MRac G5 có răng tiếp địa không và định mức bộ/tấm trước khi đặt'),
    ('Bộ đai ốc thoi M8×25 + long đen lớn (gá vào C52)', 'bộ', lambda j: f'={dlref("PV", j)}*{G["K_RN8"]}', 'hh', 'ƯL', ''),
    ('F. ĐIỂM NEO TRÊN BÈ',),
    ('Ma ní chữ U (U-Shaped Shackle) làm bích neo trên bè', 'cái', lambda j: f'=2*{dlref("LINE", j)}', 'hh', 'NEO', 'Cáp 2 nhánh chữ Y: 2 bích mỗi dây, bắt vào hộp 40x60 gia cường; mắt nhánh cáp móc thẳng vào ma ní này. Số dây theo mặt bằng neo của web (222 dây)'),
]
row = 4; stt = 0; item_rows = {}
for it in ITEMS:
    if len(it) == 1:
        th.cell(row, 2, it[0]).font = FB
        for c in range(1, 13): th.cell(row, c).fill = SEC; th.cell(row, c).border = BORDER
        row += 1; continue
    name, unit, fn, total, src, note = it
    stt += 1; r = row; item_rows[name] = r
    th.cell(r, 1, stt); th.cell(r, 2, name); th.cell(r, 4, unit); th.cell(r, 11, src); th.cell(r, 12, note)
    if total == 'cable':
        th[f'I{r}'] = f'=DuLieuBanVe!G{D["FC"]}'
    else:
        for j in range(4):
            f = fn(j)
            if '{r_prev}' in f: f = f.replace('{r_prev}', str(r - 1)).replace('{r_bars}', str(r - 3))
            th[f'{TC[j]}{r}'] = f
        th[f'I{r}'] = f'=SUM(E{r}:H{r})'
    th[f'J{r}'] = f'=ROUNDUP(I{r}*(1+{G["HH"]}/100),0)' if total == 'hh' else f'=I{r}'
    for col in 'EFGHIJ': th[f'{col}{r}'].number_format = '#,##0'
    row += 1
grid(th, 3, row - 1, 12)
for r in range(4, row):
    for col in 'AD': th[f'{col}{r}'].alignment = CENTER
for col, w in zip('ABCDEFGHIJKL', [6, 46, 23, 8, 11, 11, 11, 11, 12, 13, 8, 80]): th.column_dimensions[col].width = w
th.freeze_panes = 'E4'

# images from the earlier Huổi Vanh workbook, by item name
tpl = openpyxl.load_workbook(TEMPLATE)['TongHop']
pics = {}
for im in tpl._images:
    name = tpl.cell(im.anchor._from.row + 1, 2).value
    a = im.anchor  # shown size = the template's two-cell anchor (the picture inside column C)
    w, h = ((a.to.colOff - a._from.colOff) / 9525, (a.to.rowOff - a._from.rowOff) / 9525) if hasattr(a, 'to') and a.to.col == a._from.col else (im.width, im.height)
    pics[name] = (im._data(), w, h)
for name, r in item_rows.items():
    src = pics.get(name)
    if src:
        img = XLImage(io.BytesIO(src[0])); img.width, img.height = src[1], src[2]
        th.add_image(img, f'C{r}'); th.row_dimensions[r].height = 78
    else:
        th.cell(r, 3, '(không có hình trong tài liệu)').alignment = CENTER
        th.row_dimensions[r].height = 30

# ---- GhiChu
gh = wb.create_sheet('GhiChu')
gh['A1'] = 'GHI CHÚ, NGUỒN SỐ LIỆU VÀ CÁC ĐIỂM CẦN XÁC NHẬN'; gh['A1'].font = Font(name='Calibri', size=13, bold=True)
NOTES = [
    'Nguồn: "HOHUOIVANH.BỐ TRÍ BÈ PIN 281026.pdf" (9 tờ: CT-BE01…04 bố trí từng bè có bảng tấm pin, phao, inverter; KC-BE-01…04 khung kết cấu có bảng dầm C52 và hộp 40x60; MB-BP-01 mặt bằng tổng có bảng phao đỡ cáp), HD lắp Mibet MRac G5, tài liệu phụ kiện Mibet, datasheet G5G. Script: scripts/buildHuoiVanhRaftBOQ.py.',
    'Số tấm pin, phao, inverter, số thanh và chiều dài dầm C52, hộp 40x60 lấy NGUYÊN bảng thống kê trên bản vẽ (script kiểm tra từng số có trên đúng tờ).',
    'Số tấm pin 18.914 tấm (BÈ 1: 2.800; BÈ 2: 9.114; BÈ 3: 1.400; BÈ 4: 5.600). Từ 28/10/2026 lực gió, lực dây, cáp, cọc, đế trên web và trong thuyết minh, bảng tính, bảng vật tư neo đều tính theo số tấm này (trước đó tạm 18.354 tấm chia theo diện tích).',
    'Bản vẽ không thể hiện từng thanh thép đủ để đo: số cây cấp hàng, mối nối, bulông M8 và M10 là ƯỚC LƯỢNG theo hệ số ở GiaDinh (cùng hệ số với bảng Bản Mòng, hệ MRac G5 như nhau). Bảng KC-BE-03 (BÈ 3) ghi hộp 40x60 23 thanh dài 3.402,96 m — trung bình 148 m/thanh, dài hơn kích thước lớn nhất của bè (109,9 m); tờ KC-BE-03 vẽ tấm pin chứ không vẽ khung. Cần kiểm lại bảng này.',
    'Bản vẽ CHƯA có hộp 40x60 gia cường tai neo. Mục D bố trí cho cả 4 bè (sơ đồ GC-01, file DXF/DWG kèm theo) theo cách của bản vẽ gia cường bè 1 Huổi Vanh trước đây: cáp 2 nhánh chữ Y (Chủ đầu tư chọn 16/10/2026), 2 bích mỗi dây cách nhau 3,5 m, mỗi bích chịu K_Y × T; T là lực dây của bè trong bảng vật tư neo (engine, gió 20 m/s, chỉnh dây mỗi 1 m mực nước). Với P_KEP = 7 kN GIẢ ĐỊNH, cụm chịu lực lớn cần 2–3 thanh xếp song song — cần số liệu sức chịu móc "几" của Mibet hoặc bản vẽ gia cường của nhà cung cấp bè.',
    'Mặt bằng neo (số dây, vị trí bích) là của web, tính trên mặt bằng bè CAD 27/10/2026 (CHUGIAP-HOHUOIVANH ver 3); sheet DoiChieu_MatBang so với từng tờ CT-BE 28/10.',
    'Khối lượng thép theo tiết diện GIẢ ĐỊNH (C52 ≈ 2,36 kg/m; hộp 40×60×2,0 ≈ 3,01 kg/m) vì chưa có catalogue tiết diện của nhà sản xuất.',
    'Gá đỡ, kẹp tấm pin, bulông nối phao là ƯỚC LƯỢNG theo hệ số mỗi tấm / mỗi phao (sheet GiaDinh) — phải thay bằng BOM của nhà sản xuất khi có. Kẹp tiếp địa: hỏi Mibet trước khi đặt (có thể bỏ nếu kẹp có răng tiếp địa).',
    'Chưa gồm: cáp neo, phụ kiện đầu cáp, cọc khoan nhồi, đế neo (xem bảng thống kê vật tư neo Huổi Vanh), cáp điện DC/AC, máng cáp, tủ điện, tấm lối đi (Walkway Board) nếu có, vận chuyển.',
    'Hình minh họa ở cột C chép từ bảng vật tư bè pin Huổi Vanh trước đây (cắt từ tài liệu Mibet): chỉ để nhận dạng loại vật tư, không phải bản vẽ chế tạo.',
]
for i, t in enumerate(NOTES):
    gh.cell(3 + i, 1, i + 1); gh.cell(3 + i, 2, t).alignment = WRAP
grid(gh, 3, 2 + len(NOTES), 2)
gh.column_dimensions['A'].width = 5; gh.column_dimensions['B'].width = 150

wb._sheets = [wb['TongHop'], wb['DuLieuBanVe'], wb['GiaCuong_TaiNeo'], wb['DoiChieu_MatBang'], wb['GiaDinh'], wb['GhiChu']]
wb.active = 0
wb.save(OUT)
add_apply_flags(OUT)

# ------------------------------------------------------------------ drawing GC-01: reinforcement boxes at the anchor points
def clip_line(pts, p, d):
    """Pieces of the infinite line p + t d inside the polygon pts."""
    ts = []
    for (x1, y1), (x2, y2) in zip(pts, pts[1:] + pts[:1]):
        ex, ey = x2 - x1, y2 - y1
        den = d[0] * ey - d[1] * ex
        if abs(den) < 1e-12: continue
        t = ((x1 - p[0]) * ey - (y1 - p[1]) * ex) / den
        u = ((x1 - p[0]) * d[1] - (y1 - p[1]) * d[0]) / den
        if 0 <= u < 1: ts.append(t)
    ts.sort()
    return [(ts[i], ts[i + 1]) for i in range(0, len(ts) - 1, 2)]

dxf = ezdxf.new('R2010', setup=True)
dxf.header['$INSUNITS'] = 6
dxf.styles.add('VN', font='arial.ttf')
for lname, color, lw in (('BE_PIN', 7, 50), ('KHUNG_C52', 8, 0), ('DAY_NEO', 5, 0), ('BICH_NEO', 1, 25), ('GIA_CUONG_40x60', 1, 70), ('GHI_CHU', 7, 0)):
    dxf.layers.add(lname, color=color, lineweight=lw)
msp = dxf.modelspace()
vn = lambda v, d=0: f'{v:,.{d}f}'.replace(',', ' ').replace('.', ',').replace(' ', '.')
txt = lambda s, at, h, layer='GHI_CHU', rot=0: msp.add_text(s, height=h, rotation=rot, dxfattribs={'layer': layer, 'style': 'VN'}).set_placement(at)
nbars = {}
for name in RAFTS:
    lay = layout[name]; pts = lay['pts']
    msp.add_lwpolyline(pts, close=True, dxfattribs={'layer': 'BE_PIN'})
    # C52 every 5th beam (the frame direction; the real pitch is ~1,2 m)
    cdir = (-math.sin(lay['box_dir']), math.cos(lay['box_dir'])); bdir = (math.cos(lay['box_dir']), math.sin(lay['box_dir']))
    proj = [x * bdir[0] + y * bdir[1] for x, y in pts]
    s = min(proj) + B_C52 / 2
    while s < max(proj):
        p = (bdir[0] * s, bdir[1] * s)
        for t0, t1 in clip_line(pts, p, cdir):
            msp.add_line((p[0] + cdir[0] * t0, p[1] + cdir[1] * t0), (p[0] + cdir[0] * t1, p[1] + cdir[1] * t1), dxfattribs={'layer': 'KHUNG_C52'})
        s += 5 * B_C52
    cx = sum(x for x, _ in pts) / len(pts); cy = sum(y for _, y in pts) / len(pts)
    c = polys[name]['centroid']
    txt(f'{name} — T = {vn(T[name], 1)} kN, {lineCount[name]} dây', (c['x'] - 40, c['y']), 4.0, 'BE_PIN')
for ln in lines:
    msp.add_line((ln['xRaft'], ln['yRaft']), (ln['xAnchor'], ln['yAnchor']), dxfattribs={'layer': 'DAY_NEO'})
for name, cid, kind, codes, geo in clusters:
    n = bars_needed(len(codes), T[name]); nbars[cid] = n
    (x0, y0), (x1, y1), (sx, sy) = geo['p0'], geo['p1'], geo['step']
    for i in range(n):
        msp.add_lwpolyline([(x0 + sx * i, y0 + sy * i), (x1 + sx * i, y1 + sy * i)], dxfattribs={'layer': 'GIA_CUONG_40x60', 'const_width': 0.3})
    for (bx, by) in geo['cleats']:
        msp.add_circle((bx, by), 0.6, dxfattribs={'layer': 'BICH_NEO'})
    ang = math.degrees(math.atan2(y1 - y0, x1 - x0))
    if ang > 90: ang -= 180
    if ang < -90: ang += 180
    mx, my = (x0 + x1) / 2 + sx * (n + 0.5), (y0 + y1) / 2 + sy * (n + 0.5)
    txt(f'{cid}: {n}×{L_GC:.0f} m', (mx, my), 1.2, 'GIA_CUONG_40x60', ang)
xs = [x for n in RAFTS for x, _ in layout[n]['pts']]; ys = [y for n in RAFTS for _, y in layout[n]['pts']]
X0, Y0 = max(xs) + 60, max(ys)
txt('SƠ ĐỒ BỐ TRÍ HỘP 40x60 GIA CƯỜNG TẠI CÁC VỊ TRÍ NEO — HỒ HUỔI VANH, 4 BÈ (GC-01)', (X0, Y0), 6.0)
txt('Mặt bằng bè: CAD 27/10/2026; khung kết cấu: KC-BE-01…04 (28/10/2026); mặt bằng neo: web (222 dây, cáp 2 nhánh chữ Y). Đơn vị: m.', (X0, Y0 - 10), 3.0)
y = Y0 - 22
for n in RAFTS:
    cl = [c for c in clusters if c[0] == n]
    nb = sum(nbars[c[1]] for c in cl)
    txt(f'{n}: {len(cl)} cụm, {nb} thanh × {L_GC:.0f} m = {vn(nb * L_GC)} m hộp 40x60 gia cường; T = {vn(T[n], 1)} kN; dầm C52 {"ngang" if C52_HORIZONTAL[n] else "đứng"} trên tờ KC-BE', (X0, y), 3.0); y -= 6
tot = sum(nbars.values())
txt(f'TỔNG: {len(clusters)} cụm, {tot} thanh, {vn(tot * L_GC)} m', (X0, y), 3.5); y -= 10
for t in ['GHI CHÚ:',
          '1. Mỗi dây neo có 2 bích (cáp 2 nhánh chữ Y) cách nhau 3,5 m, vòng tròn đỏ. Mã cụm và số thanh: sheet GiaCuong_TaiNeo của bảng vật tư bè pin.',
          '2. Bích trên cạnh song song với hộp 40x60: các bích trong đoạn 14 m dùng chung thanh dọc cạnh, cách mép 1 hàng phao.',
          '   Bích trên cạnh vuông góc: mỗi bích một thanh từ mép chạy vào trong, cùng phương hộp 40x60 kết cấu.',
          '3. Hộp kẹp vào MỌI dầm C52 nó cắt qua bằng móc "几" (2 đai ốc thoi M10×30 + 1 bulông M10×120).',
          f'4. Số thanh = MAX(1; số bích × 0,7 T × 1,2 / (13 giao C52 × 7 kN)); sức chịu móc 7 kN là GIẢ ĐỊNH, chờ số liệu Mibet.',
          '5. Nét xám mảnh: phương dầm C52 (vẽ 1/5 số dầm). Nét xanh: dây neo. Đây là SƠ ĐỒ BỐ TRÍ, không phải bản vẽ chế tạo:',
          '   vị trí chính xác theo hàng phao và chi tiết liên kết do nhà cung cấp bè lập và xác nhận.']:
    txt(t, (X0, y), 2.6); y -= 5
dxf.saveas(OUT_DXF)

print(json.dumps({'out': OUT, 'dxf': OUT_DXF, 'clusters': len(clusters), 'bars': tot,
                  'byRaft': {n: [sum(1 for c in clusters if c[0] == n), sum(nbars[c[1]] for c in clusters if c[0] == n)] for n in RAFTS},
                  'outline': {n: [round(layout[n]['area']), polys[n]['area_m2'], f"{layout[n]['w']:.1f}x{layout[n]['h']:.1f}", f"{layout[n]['Wu']:.1f}x{layout[n]['Wv']:.1f}"] for n in RAFTS}}, ensure_ascii=False))
