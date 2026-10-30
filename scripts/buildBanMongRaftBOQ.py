"""
Bảng thống kê vật tư BÈ PIN hồ Bản Mòng, mặt bằng 6 bè (chốt 30/10/2026), cùng khuôn với Huổi Vanh
(scripts/buildHuoiVanhRaftBOQ.py), kèm sơ đồ tổng hộp 40x60 gia cường tại các vị trí neo GC-01:

  python scripts/buildBanMongRaftBOQ.py <BO TRI BE PIN HO BAN MONG.pdf> <bảng vật tư neo .xlsx> <mẫu hình .xlsx> <out.xlsx> <out.dxf> <reinforcement.json> <thư mục kc1…kc6.dxf>

- Số tấm pin, dải pin, phao từng loại, inverter: lấy NGUYÊN bảng Revit trên các tờ CT-BP-01…06 của "BO TRI BE PIN HO BAN MONG.pdf"
  (Chủ đầu tư 30/10/2026; script kiểm tổng từng bảng); phao đỡ cáp ở tờ MB-B01-01.
- Dầm C52, hộp 40x60: ĐO từng thanh trên 6 tờ khung kết cấu Revit "CHUGIAP-HOBANMONG ver 2-Sheet - CT-KC-BE-01…06 - BỐ TRÍ KHUNG
  KẾT CẤU BÈ PIN 0n.dxf" (reinforcement.json của scripts/buildHuoiVanhRaftReinforcement.py --lake=ban-mong); bảng "THỐNG KÊ KHUNG
  THÉP" trên tờ khung để đối chiếu (sheet DoiChieu_KhungThep, số của bảng ghi ở KC_TABLE).
- Đường bao từng bè đo trên ảnh tờ CT-BP (sheet DoiChieu_MatBang) để so với mặt bằng mà web dùng tính neo.
- Hộp 40x60 gia cường tai neo: theo bản vẽ khung có gia cường (lớp GC-*), lực dây T của bè theo sheet KiemTra_Nhanh của bảng
  vật tư neo; <out.dxf> là sơ đồ tổng GC-01 trên mặt bằng neo (mét).
Hình minh họa ở cột C chép từ <mẫu hình .xlsx> (bảng vật tư bè pin Huổi Vanh 9 bè) theo tên hạng mục.
"""
import io, json, math, re, sys
from pathlib import Path

import cv2
import ezdxf
import numpy as np
import openpyxl
import pymupdf
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as COL
from shapely.geometry import box as sbox
from shapely.ops import unary_union

from xlsx_apply_flags import add_apply_flags

ROOT = Path(__file__).resolve().parent.parent
PDF, MOORING_XLSX, TEMPLATE, OUT, OUT_DXF, REINF, KC_DIR = sys.argv[1:8]
RJ = json.load(open(REINF, encoding='utf-8'))['rafts']
RAFTS = [f'BÈ {i}' for i in range(1, 7)]
NR = len(RAFTS)
assert set(RAFTS) == set(RJ), RJ.keys()


def kc_table(path):
    """Table "THỐNG KÊ KHUNG THÉP" of a frame sheet kc<i>.dxf (MTEXT on G-ANNO-SCHD, read from the raw DXF stream):
    ((C52 bars, m), (box 40x60 bars, m))."""
    rows, ent, layer, xy, text = {}, None, None, [0, 0], ''
    with open(path, encoding='utf-8', errors='replace') as f:
        it = iter(f)
        for code in it:
            code = code.strip(); val = next(it).rstrip('\r\n')
            if code == '0':
                if ent == 'MTEXT' and layer == 'G-ANNO-SCHD':
                    t = re.sub(r'\\U\+([0-9A-Fa-f]{4})', lambda m: chr(int(m.group(1), 16)), text).split(';')[-1].strip('{} ')
                    rows.setdefault(round(xy[1], 1), []).append((xy[0], t))
                ent, layer, text = val, None, ''
            elif code == '8': layer = val
            elif code == '10': xy[0] = float(val)
            elif code == '20': xy[1] = float(val)
            elif code in ('1', '3'): text += val
    out = {}
    for cells in rows.values():
        cells = [t for _, t in sorted(cells)]
        if cells and ('C52' in cells[0] or '40x60' in cells[0]):
            nums = [c for c in cells[1:] if re.fullmatch(r'[\d.]+', c)]
            out['C52' if 'C52' in cells[0] else 'BOX'] = (int(nums[0]), float(nums[1]))
    return out['C52'], out['BOX']


# table "THỐNG KÊ KHUNG THÉP" of the frame sheets CT-KC-BE-01…06: C52 (bars, m), box 40x60 (bars, m)
KC_TABLE = {n: kc_table(Path(KC_DIR) / f'kc{i}.dxf') for i, n in enumerate(RAFTS, start=1)}

# ------------------------------------------------------------------ tables of the sheets CT-BP (Revit)
doc = pymupdf.open(PDF)


def after(ls, *keys, n=2, start=0):
    """The n lines after the first run of consecutive lines `keys` at or after `start`, and where the run starts."""
    for i in range(start, len(ls) - len(keys)):
        if all(ls[i + k] == key for k, key in enumerate(keys)):
            return ls[i + len(keys):i + len(keys) + n], i
    raise KeyError(keys)


num = lambda s: int(s.replace('.', '').replace(',', ''))
SHEETS = {}
for i, name in enumerate(RAFTS, start=1):
    page = next(k for k, p in enumerate(doc) if f'CT-BP-0{i}' in p.get_text())
    ls = [l.strip() for l in doc[page].get_text().splitlines()]
    assert f'BỐTRÍ BÈ PIN 0{i}' in ls, name
    scale = int(re.search(r'1 : (\d+)', doc[page].get_text()).group(1))
    (fs_r, fs_n), _ = after(ls, 'Dai phao do pin', 'DẢI PHAO ĐỠ PIN')
    (wi_r, wi_n), _ = after(ls, 'Dải phao đi lại kèm', 'đỡ invester', 'DẢI PHAO ĐI LẠI')
    (wa_r, wa_n), _ = after(ls, 'Dai phaodi lai', 'DẢI PHAO ĐI LẠI')
    _, k = after(ls, 'SỐ PHAO NỔI THEO BÈ', n=0)
    tot = next(ls[j + 1] for j in range(k, len(ls)) if ls[j].startswith('Grand total'))
    (pv_r, pv_n), k = after(ls, 'DẢI PIN MẶT TRỜI')
    strings = next(ls[j + 2] for j in range(k, len(ls)) if ls[j].startswith('Grand total'))
    (inv,), _ = after(ls, 'SG -510HX', n=1)
    s_ = dict(page=page + 1, scale=scale, pv=num(pv_n), dpv=num(pv_r), strings=num(strings), inv=num(inv),
              fs=(num(fs_n), num(fs_r)), fwi=(num(wi_n), num(wi_r)), fwa=(num(wa_n), num(wa_r)), total=num(tot))
    assert s_['fs'][0] + s_['fwi'][0] + s_['fwa'][0] == s_['total'], (name, s_)
    SHEETS[name] = s_
mb = next(k for k, p in enumerate(doc) if 'MB-B01-01' in p.get_text())
(cr, cn), _ = after([l.strip() for l in doc[mb].get_text().splitlines()], 'ĐỠ CÁP')
CABLE_FLOATS = (num(cr), num(cn))  # rows, floats (whole lake)


def sheet_outline(page, scale, dpi=150):
    """Area (m²) of the raft drawn on a CT-BP sheet (filled, tables and dimensions dropped)."""
    p = doc[page - 1]
    clip = pymupdf.Rect(28, 7, 1042, 835)  # drawing area, left of the title block
    pix = p.get_pixmap(dpi=dpi, clip=clip, colorspace=pymupdf.csGRAY)
    a = np.frombuffer(pix.samples, np.uint8).reshape(pix.h, pix.w)
    mpp = 25.4 / 72 * scale / 1000 / (dpi / 72)
    b = (a < 200).astype(np.uint8)
    b = cv2.morphologyEx(b, cv2.MORPH_CLOSE, np.ones((max(3, int(1.5 / mpp)),) * 2, np.uint8))
    b = cv2.morphologyEx(b, cv2.MORPH_OPEN, np.ones((max(5, int(6 / mpp)),) * 2, np.uint8))
    n, lab, st_, _ = cv2.connectedComponentsWithStats(b, 8)
    m = (lab == 1 + int(np.argmax(st_[1:, 4]))).astype(np.uint8)
    cnt = max(cv2.findContours(m, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)[0], key=cv2.contourArea)
    return cv2.contourArea(cnt) * mpp * mpp

data = ROOT / 'src' / 'data' / 'banMong'
polys = {r['name']: r for r in json.load(open(data / 'raftPolygons.json', encoding='utf-8'))}
catalogue = {r['name']: r for r in json.load(open(data / 'raftCatalogue.json', encoding='utf-8'))}
lines = json.load(open(data / 'coordinates.json', encoding='utf-8'))

mw = openpyxl.load_workbook(MOORING_XLSX, data_only=True)['KiemTra_Nhanh']
T = {r[0]: r[1] for r in mw.iter_rows(min_row=4, values_only=True) if r[0] in RAFTS}
assert set(T) == set(RAFTS), T
for n in RAFTS: assert abs(RJ[n]['T_kN'] - T[n]) < 0.05, (n, RJ[n]['T_kN'], T[n])

L_GC, P_KEP, K_Y, GAMMA = 15.0, 7.0, 0.7, 1.2
B_C52 = round(sum(RJ[n]['c52Pitch_m'] for n in RAFTS) / NR, 3)
clusters = [(n, c['id'], c['kind'], c['cleats'], c) for n in RAFTS for c in RJ[n]['clusters']]
C52_CAD = {n: RJ[n]['c52_m'] for n in RAFTS}
BOX_CAD = {n: RJ[n]['box_m'] for n in RAFTS}
lineCount = {n: catalogue[n]['cableCount'] for n in RAFTS}
assert sum(lineCount.values()) == len(lines), (sum(lineCount.values()), len(lines))
for n in RAFTS: assert catalogue[n]['solarPanelCount'] == SHEETS[n]['pv'], (n, catalogue[n]['solarPanelCount'], SHEETS[n]['pv'])


def frame_area(i):
    """Footprint of the frame (C52, boxes, float strips) of the drawing, m²."""
    M = json.load(open(Path(REINF).parent / f'kc{i}.members.json', encoding='utf-8'))['members']
    return unary_union([sbox(r[1] / 1000, r[2] / 1000, r[3] / 1000, r[4] / 1000).buffer(1.5) for r in M]).buffer(-1.5).area


# ------------------------------------------------------------------ workbook
wb = openpyxl.Workbook()
thin = Side(style='thin'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='F2F2F2'); SEC = PatternFill('solid', fgColor='D9D9D9'); INPUT = PatternFill('solid', fgColor='FFFF99')
F = Font(name='Calibri', size=12); FB = Font(name='Calibri', size=12, bold=True)
WRAP = Alignment(wrap_text=True, vertical='center'); CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
RC = [COL(3 + j) for j in range(NR)]          # raft columns of DuLieuBanVe (C…H)
D_TOT, D_SRC = COL(3 + NR), COL(4 + NR)       # I, J
TC = [COL(5 + j) for j in range(NR)]          # raft columns of TongHop (E…J)
T_TOT, T_HH, T_SRC, T_NOTE = (COL(5 + NR + k) for k in range(4))  # K, L, M, N


def grid(ws, r0, r1, c1):
    for row in ws.iter_rows(min_row=r0, max_row=r1, max_col=c1):
        for c in row:
            c.border = BORDER
            if not c.font.b: c.font = F
            if c.alignment.horizontal is None: c.alignment = WRAP


# ---- GiaDinh
gd = wb.active; gd.title = 'GiaDinh'
GD = [
    ('L_C52', 'Chiều dài cây dầm C52 khi cấp hàng', 6, 'm', 'GIẢ ĐỊNH. HD lắp Mibet G5 dùng đoạn L2990 / L3750 — nếu cấp đoạn ~3 m thì số mối nối gần gấp đôi'),
    ('L_H', 'Chiều dài cây hộp 40x60', 6, 'm', 'GIẢ ĐỊNH, cây thép hộp thương mại 6 m'),
    ('M_C52', 'Khối lượng dầm C52', 2.36, 'kg/m', 'GIẢ ĐỊNH tiết diện C 52×41×12×2,0 (chưa có catalogue nhà sản xuất)'),
    ('M_H', 'Khối lượng hộp 40x60', 3.01, 'kg/m', 'GIẢ ĐỊNH 40×60×2,0'),
    ('K_FS', 'Gá trước (Front Support) mỗi tấm pin', 2, 'cái/tấm', 'ƯỚC LƯỢNG theo HD lắp G5'),
    ('K_BS', 'Gá sau (Back Support) mỗi tấm pin', 2, 'cái/tấm', 'ƯỚC LƯỢNG'),
    ('K_EC', 'Bộ kẹp biên tấm pin (End Clamp Kit) mỗi tấm', 2, 'bộ/tấm', 'ƯỚC LƯỢNG'),
    ('K_GC', 'Bộ kẹp tiếp địa mỗi tấm', 2, 'bộ/tấm', 'ƯỚC LƯỢNG; hỏi Mibet: kẹp biên/kẹp giữa có răng tiếp địa thì bỏ được hạng mục này, hoặc 1 bộ/tấm'),
    ('K_RN8', 'Bộ đai ốc thoi M8×25 bắt gá vào C52', 4, 'bộ/tấm', 'ƯỚC LƯỢNG = 1 bộ mỗi gá'),
    ('K_FB', 'Bộ bulông nhựa nối phao', 2, 'bộ/phao', 'ƯỚC LƯỢNG (HD: các phao nối với nhau bằng bulông + đai ốc nhựa)'),
    ('HH', 'Hao hụt phụ kiện, bulông', 3, '%', 'GIẢ ĐỊNH, chỉ cộng ở cột "Tổng có hao hụt"'),
    ('L_GC', 'Chiều dài 1 thanh hộp 40x60 gia cường tai neo', L_GC, 'm', 'Như Huổi Vanh (bản vẽ gia cường bè 1 Huổi Vanh 9 bè: 13,4–15,5 m)'),
    ('P_KEP', 'Sức chịu 1 móc "几" (2 đai ốc thoi M10×30 trong rãnh C52)', P_KEP, 'kN', 'GIẢ ĐỊNH — chưa có số liệu nhà sản xuất; quyết định số thanh xếp song song ở cụm chịu lực lớn. Bản vẽ vẽ số thanh theo giá trị mặc định này'),
    ('GAMMA', 'Hệ số tải trọng cho lực dây', GAMMA, '', 'Như kiểm tra cọc (pileBendingLoadFactor)'),
    ('K_Y', 'Phần lực dây vào MỘT bích (cáp 2 nhánh chữ Y)', K_Y, '× T', 'Chia đều ở góc 60°: 0,58 T; lấy 0,7 T vì 2 nhánh có thể dài lệch nhau (cùng giả định với bảng vật tư neo, sheet KiemTra_Nhanh)'),
    ('K_XF', 'Số chỗ dầm C52 vượt qua một hàng phao, trên 1 m dầm', 0.5, 'điểm/m', 'ĐO trên bản vẽ CAD khung: hàng phao đỡ pin cách nhau ~2,0 m'),
    ('B_C52', 'Khoảng cách trung bình giữa các dầm C52 (theo phương hộp)', B_C52, 'm', 'ĐO trên bản vẽ CAD khung 6 bè. Chỉ dùng ước lượng bulông M10×120 hộp kết cấu – dầm (1 bộ mỗi giao điểm); hộp gia cường đếm giao điểm trên bản vẽ'),
]
gd['A1'] = 'THÔNG SỐ GIẢ ĐỊNH — ô tô vàng là ô NHẬP, sửa là cả file tự tính lại (bản vẽ gia cường vẽ theo giá trị mặc định)'; gd['A1'].font = Font(name='Calibri', size=13, bold=True)
for j, h in enumerate(['Mã', 'Thông số', 'Giá trị', 'Đơn vị', 'Ghi chú']):
    c = gd.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
G = {}
for i, (code, label, val, unit, note) in enumerate(GD):
    r = 4 + i; G[code] = f'GiaDinh!$C${r}'
    for j, v in enumerate([code, label, val, unit, note]):
        gd.cell(r, j + 1, v)
    gd.cell(r, 3).fill = INPUT
grid(gd, 3, 3 + len(GD), 5)
for col, w in zip('ABCDE', [9, 52, 10, 10, 110]): gd.column_dimensions[col].width = w

# ---- DuLieuBanVe
dl = wb.create_sheet('DuLieuBanVe', 0)
dl['A1'] = 'SỐ LIỆU LẤY TỪ BẢN VẼ "BO TRI BE PIN HO BAN MONG" (bảng Revit trên các tờ CT-BP-01…06, MB-B01-01) VÀ BẢN VẼ CAD KHUNG "CT-KC-BE-01…06" (30/10/2026)'; dl['A1'].font = Font(name='Calibri', size=13, bold=True)
for j, h in enumerate(['Mã', 'Hạng mục'] + RAFTS + ['TỔNG', 'Nguồn']):
    c = dl.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
DL = [
    ('PV', 'Tấm pin (tấm)', lambda n: SHEETS[n]['pv'], 'Bảng "THỐNG KÊ TẤM PIN MẶT TRỜI" (tờ CT-BP), cột Số tấm pin'),
    ('DPV', 'Dải pin (hàng)', lambda n: SHEETS[n]['dpv'], 'Như trên, cột SỐ LƯỢNG HÀNG'),
    ('STR', 'Số string', lambda n: SHEETS[n]['strings'], 'Như trên, cột So string'),
    ('I510', 'Inverter SG-510HX (bộ)', lambda n: SHEETS[n]['inv'], 'Bảng "INVESTER" (tờ CT-BP)'),
    ('FS', 'Phao đỡ pin — Main Floater (cái)', lambda n: SHEETS[n]['fs'][0], 'Bảng "SỐ PHAO NỔI THEO BÈ", dòng "Dai phao do pin"'),
    ('FSR', 'Số hàng phao đỡ pin', lambda n: SHEETS[n]['fs'][1], 'Như trên, cột SỐ LƯỢNG HÀNG'),
    ('FWA', 'Phao đi lại kèm đỡ inverter (cái)', lambda n: SHEETS[n]['fwi'][0], 'Dòng "Dải phao đi lại kèm đỡ invester"'),
    ('FWB', 'Phao đi lại (cái)', lambda n: SHEETS[n]['fwa'][0], 'Dòng "Dai phaodi lai"'),
    ('FWR', 'Số hàng phao đi lại', lambda n: SHEETS[n]['fwi'][1] + SHEETS[n]['fwa'][1], 'Cộng cột SỐ LƯỢNG HÀNG hai dòng phao đi lại'),
    ('FT', 'Tổng phao (cái)', None, 'Khớp "Grand total" của bảng phao từng tờ'),
    ('C52N', 'Dầm C52 — số thanh (ĐO CAD)', lambda n: len(C52_CAD[n]), 'Đếm trên bản vẽ CAD khung (sheet C52_TungThanh)'),
    ('C52L', 'Dầm C52 — chiều dài (m, ĐO CAD)', lambda n: round(sum(C52_CAD[n]), 2), 'Cộng từng thanh trên bản vẽ CAD'),
    ('HN', 'Hộp 40x60 kết cấu — số thanh (ĐO CAD)', lambda n: len(BOX_CAD[n]), 'Đếm trên bản vẽ CAD (sheet Hop40x60_TungThanh)'),
    ('HL', 'Hộp 40x60 kết cấu — chiều dài (m, ĐO CAD)', lambda n: round(sum(BOX_CAD[n]), 2), 'Cộng từng thanh trên bản vẽ CAD'),
    ('C52N_BV', 'Dầm C52 — số thanh (bảng tờ CT-KC-BE)', lambda n: KC_TABLE[n][0][0], 'Bảng "THỐNG KÊ KHUNG THÉP" trên tờ khung — để đối chiếu'),
    ('C52L_BV', 'Dầm C52 — chiều dài (m, bảng tờ CT-KC-BE)', lambda n: KC_TABLE[n][0][1], 'Như trên'),
    ('HN_BV', 'Hộp 40x60 — số thanh (bảng tờ CT-KC-BE)', lambda n: KC_TABLE[n][1][0], 'Như trên'),
    ('HL_BV', 'Hộp 40x60 — chiều dài (m, bảng tờ CT-KC-BE)', lambda n: KC_TABLE[n][1][1], 'Như trên'),
    ('LINE', 'Số dây neo của bè (tuyến cáp)', lambda n: lineCount[n], 'Mặt bằng neo trên web (ưu tiên bờ, 80 m)'),
]
D = {}
for i, (code, label, fn, src) in enumerate(DL):
    r = 4 + i; D[code] = r
    dl.cell(r, 1, code); dl.cell(r, 2, label); dl[f'{D_SRC}{r}'] = src
    for j, n in enumerate(RAFTS):
        col = RC[j]
        dl[f'{col}{r}'] = f'={col}{D["FS"]}+{col}{D["FWA"]}+{col}{D["FWB"]}' if code == 'FT' else fn(n)
        dl[f'{col}{r}'].number_format = '#,##0.00' if code in ('C52L', 'HL', 'C52L_BV', 'HL_BV') else '#,##0'
    dl[f'{D_TOT}{r}'] = f'=SUM({RC[0]}{r}:{RC[-1]}{r})'; dl[f'{D_TOT}{r}'].number_format = dl[f'{RC[0]}{r}'].number_format
r = 4 + len(DL); D['FC'] = r
dl.cell(r, 1, 'FC'); dl.cell(r, 2, 'Phao đỡ cáp điện giữa các bè (cái, cả hồ)'); dl[f'{D_TOT}{r}'] = CABLE_FLOATS[1]
dl[f'{D_SRC}{r}'] = f'Tờ MB-B01-01, bảng "SỐ PHAO ĐỠ CÁP": {CABLE_FLOATS[0]} hàng, {CABLE_FLOATS[1]} phao (không chia theo bè)'
LAST_DL = r
grid(dl, 3, LAST_DL, 4 + NR)
for j, w in enumerate([8, 44] + [11] * NR + [12, 90]): dl.column_dimensions[COL(j + 1)].width = w

# ---- GiaCuong_TaiNeo
gc = wb.create_sheet('GiaCuong_TaiNeo')
gc['A1'] = 'HỘP 40x60 GIA CƯỜNG TẠI VỊ TRÍ NEO — VẼ TRÊN BẢN VẼ KHUNG KẾT CẤU "CT-KC-BE-01…06" (lớp GC-*), CÁP 2 NHÁNH CHỮ Y (2 bích mỗi dây, cách nhau 3,5 m)'
gc['A2'] = ('Mỗi thanh nằm trên một hàng phao, theo phương hộp 40x60 kết cấu (vuông góc dầm C52), kẹp móc "几" vào MỌI dầm C52 nó cắt qua — số dầm cắt qua ĐẾM trên bản vẽ khung. '
            'Bích trên cạnh song song với hộp: các bích trong 14 m dùng chung thanh dọc cạnh (hàng phao thứ 2 tính từ mép); bích trên cạnh vuông góc: mỗi bích một thanh từ dầm C52 ngoài cùng vào trong. '
            'Số thanh cần = MAX(1; ROUNDUP(số bích × K_Y × T × γ / (dầm C52 cắt qua × P_KEP))); cột "Số thanh vẽ" là số thanh trên bản vẽ (theo giá trị mặc định ở GiaDinh) — sửa P_KEP mà cột Kiểm tra báo THIẾU thì phải vẽ thêm thanh.')
for a in ('A1', 'A2'): gc[a].font = FB if a == 'A1' else F
gc.merge_cells('A2:S2'); gc['A2'].alignment = WRAP; gc.row_dimensions[2].height = 62
HDR = ['Bè', 'Mã cụm (trên bản vẽ)', 'Cách bố trí', 'Bích trong cụm (mã dây + a/b)', 'Số bích', 'T dây (kN, gió 20 m/s)', 'Lực thiết kế cụm = số bích × K_Y × T × γ (kN)', 'Dầm C52 cắt qua — thanh ít nhất (ĐO CAD)',
       'Sức chịu 1 thanh (kN)', 'Số thanh cần', 'Số thanh vẽ', 'Kiểm tra', 'Tổng dài (m, ĐO CAD)', 'Số cây cấp hàng', 'Số mối nối', 'Móc "几" (cái, ĐO CAD)', 'Đai ốc thoi M10×30 (bộ)', 'Bulông M10×120 (bộ)', 'Khối lượng (kg)']
for j, h in enumerate(HDR):
    c = gc.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
gc.row_dimensions[3].height = 62
for i, (name, cid, kind, codes, c) in enumerate(clusters):
    r = 4 + i
    bars = c['bars']
    pieces = '+'.join(f'ROUNDUP({b["L_m"]:.3f}/{G["L_H"]}-0.0005,0)' for b in bars)
    vals = [name, cid, kind, ', '.join(codes), len(codes), T[name], f'=E{r}*F{r}*{G["K_Y"]}*{G["GAMMA"]}', min(b['nC52'] for b in bars),
            f'=H{r}*{G["P_KEP"]}', f'=MAX(1,ROUNDUP(G{r}/I{r}-0.0005,0))', len(bars), f'=IF(K{r}>=J{r},"ĐỦ","THIẾU")', round(sum(b['L_m'] for b in bars), 3),
            f'={pieces}', f'=N{r}-K{r}', sum(b['nC52'] for b in bars), f'=2*P{r}', f'=P{r}', f'=M{r}*{G["M_H"]}']
    for j, v in enumerate(vals): gc.cell(r, j + 1, v)
    gc[f'G{r}'].number_format = '0.0'; gc[f'I{r}'].number_format = '0.0'; gc[f'M{r}'].number_format = '0.0'; gc[f'S{r}'].number_format = '#,##0.0'
GC_LAST = 3 + len(clusters)
grid(gc, 3, GC_LAST, 19)
for col, w in zip('ABCDEFGHIJKLMNOPQRS', [8, 10, 18, 34, 7, 10, 14, 12, 10, 9, 9, 9, 10, 9, 8, 10, 10, 10, 11]): gc.column_dimensions[col].width = w

# ---- DoiChieu_MatBang
dc = wb.create_sheet('DoiChieu_MatBang')
dc['A1'] = 'ĐỐI CHIẾU TỜ CT-BP VÀ KHUNG CT-KC-BE (30/10/2026) VỚI MẶT BẰNG BÈ MÀ WEB DÙNG TÍNH NEO (bản vẽ CAD bố trí bè 30/10/2026)'; dc['A1'].font = Font(name='Calibri', size=13, bold=True)
dc['A2'] = ('Diện tích tờ CT-BP đo trên ảnh tờ (tỷ lệ in trên tờ, gồm cả mép phao); diện tích khung = phần phủ của dầm C52, hộp 40x60 và dải phao trên bản vẽ khung. '
            'Bản vẽ khung là mặt bằng neo xoay quanh gốc tọa độ: góc xoay và độ lệch khi khớp đường bao vào mép dải phao.')
dc.merge_cells('A2:K2'); dc['A2'].alignment = WRAP; dc.row_dimensions[2].height = 34
H2 = ['Bè', 'Diện tích khung CAD (m²)', 'Diện tích trên web (m²)', 'Lệch diện tích', 'Tấm pin (tờ CT-BP = web)', 'Dầm C52 TB 1 thanh (m)', 'Hộp 40x60 TB 1 thanh (m)', 'Góc xoay khung (°)', 'Lệch khớp TB (m)', 'Số dây neo', 'Diện tích tờ CT-BP (m²)']
for j, h in enumerate(H2):
    c = dc.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
dc.row_dimensions[3].height = 48
for i, n in enumerate(RAFTS):
    r = 4 + i
    row = [n, round(frame_area(i + 1)), polys[n]['area_m2'], f'=B{r}/C{r}-1', catalogue[n]['solarPanelCount'],
           f"=DuLieuBanVe!{RC[i]}{D['C52L']}/DuLieuBanVe!{RC[i]}{D['C52N']}", f"=DuLieuBanVe!{RC[i]}{D['HL']}/DuLieuBanVe!{RC[i]}{D['HN']}",
           round(RJ[n]['fit']['rot_deg'], 2), round(RJ[n]['fit']['meanDist_m'], 2), lineCount[n], round(sheet_outline(SHEETS[n]['page'], SHEETS[n]['scale']))]
    for j, v in enumerate(row): dc.cell(r, j + 1, v)
    dc[f'D{r}'].number_format = '+0.0%;-0.0%'
    for col in 'FG': dc[f'{col}{r}'].number_format = '0.0'
    for col in 'BCEK': dc[f'{col}{r}'].number_format = '#,##0'
r = 4 + NR
dc.cell(r, 1, 'TỔNG')
for col in 'BCEJK': dc[f'{col}{r}'] = f'=SUM({col}4:{col}{r - 1})'; dc[f'{col}{r}'].number_format = '#,##0'
dc[f'D{r}'] = f'=B{r}/C{r}-1'; dc[f'D{r}'].number_format = '+0.0%;-0.0%'
grid(dc, 3, r, 11)
for c in dc[r]: c.font = FB
for col, w in zip('ABCDEFGHIJK', [8, 13, 12, 10, 12, 11, 11, 10, 10, 9, 10]): dc.column_dimensions[col].width = w

# ---- C52_TungThanh, Hop40x60_TungThanh
BAR_LAST = {}
for sheet, data_, stock, title in (('C52_TungThanh', C52_CAD, 'L_C52', 'DẦM C52'), ('Hop40x60_TungThanh', BOX_CAD, 'L_H', 'HỘP 40x60 KẾT CẤU')):
    ws = wb.create_sheet(sheet)
    ws['A1'] = f'{title} — TỪNG THANH ĐO TRÊN BẢN VẼ CAD "CT-KC-BE-01…06" (khung bao của từng thanh Revit, m)'; ws['A1'].font = FB
    for j, h in enumerate(['Bè', 'STT', 'Chiều dài thanh (m)', f'Số cây cấp hàng = ROUNDUP(L / {stock})', 'Số mối nối = số cây − 1']):
        c = ws.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
    r = 4
    for n in RAFTS:
        for k, L in enumerate(data_[n], start=1):
            ws.cell(r, 1, n); ws.cell(r, 2, k); ws.cell(r, 3, L)
            ws.cell(r, 4, f'=ROUNDUP(C{r}/{G[stock]}-0.0005,0)'); ws.cell(r, 5, f'=D{r}-1'); ws.cell(r, 3).number_format = '0.000'
            r += 1
    BAR_LAST[sheet] = r - 1
    for col, w in zip('ABCDE', [8, 6, 14, 22, 18]): ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A4'


def barsum(sheet, col, n): return f"SUMIF('{sheet}'!$A$4:$A${BAR_LAST[sheet]},\"{n}\",'{sheet}'!${col}$4:${col}${BAR_LAST[sheet]})"  # quoted: "C52_…" would read as cell C52


# ---- TongHop
th = wb.create_sheet('TongHop', 0)
th['A1'] = 'BẢNG TỔNG HỢP VẬT TƯ BÈ PIN — 6 BÈ HỒ BẢN MÒNG (khung kết cấu "CT-KC-BE-01…06" 30/10/2026, HD lắp Mibet MRac G5)'
th['A2'] = ('Nguồn: BV = lấy đúng bảng Revit trên tờ CT-BP; CAD = đếm / đo trên bản vẽ khung kết cấu; ƯL = ước lượng theo thông số sheet GiaDinh '
            '(phải đối chiếu BOM nhà sản xuất); NEO = theo mặt bằng neo trên web. Mục D vẽ trên bản vẽ khung có gia cường (lớp GC-*).')
th.merge_cells(f'A1:{T_NOTE}1'); th.merge_cells(f'A2:{T_NOTE}2')
th['A1'].font = Font(name='Calibri', size=13, bold=True); th['A2'].alignment = WRAP; th.row_dimensions[1].height = 17; th.row_dimensions[2].height = 34
for j, h in enumerate(['STT', 'Hạng mục', 'Hình minh họa (tài liệu Mibet)', 'Đơn vị'] + RAFTS + ['TỔNG', 'Tổng có hao hụt', 'Nguồn', 'Ghi chú']):
    c = th.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
th.row_dimensions[3].height = 31


def dlref(code, j): return f'DuLieuBanVe!{RC[j]}{D[code]}'
def gcsum(col, n): return f'SUMIF(GiaCuong_TaiNeo!$A$4:$A${GC_LAST},"{n}",GiaCuong_TaiNeo!${col}$4:${col}${GC_LAST})'


NPV = sum(catalogue[n]['solarPanelCount'] for n in RAFTS)
NINV = sum(SHEETS[n]['inv'] for n in RAFTS)
ITEMS = [
    ('A. PIN, INVERTER, PHAO NỔI',),
    ('Tấm pin mặt trời', 'tấm', lambda j: '=' + dlref('PV', j), 'same', 'BV', f'{NPV:,} tấm theo bảng từng tờ CT-BP (30/10/2026); lực gió và lực dây trên web tính theo số tấm từng bè này'.replace(',', '.')),
    ('Inverter SG-510HX', 'bộ', lambda j: '=' + dlref('I510', j), 'same', 'BV', f'{NINV} inverter; bản vẽ không có SG350HX'),
    ('Phao đỡ pin (Main Floater G5G)', 'cái', lambda j: '=' + dlref('FS', j), 'same', 'BV', ''),
    ('Phao đi lại (Walkway Floater) — nhóm 1', 'cái', lambda j: '=' + dlref('FWA', j), 'same', 'BV', 'Dải phao đi lại kèm đỡ inverter. Bản vẽ chia phao đi lại thành 2 dòng; cần nhà sản xuất xác nhận là 1 hay 2 loại phao'),
    ('Phao đi lại (Walkway Floater) — nhóm 2', 'cái', lambda j: '=' + dlref('FWB', j), 'same', 'BV', 'Dải phao đi lại'),
    ('Tổng phao', 'cái', lambda j: '=' + dlref('FT', j), 'same', 'BV', ''),
    ('Phao đỡ cáp điện giữa các bè', 'cái', None, 'cable', 'BV', f'Tờ MB-B01-01: {CABLE_FLOATS[0]:,} hàng × 2 = {CABLE_FLOATS[1]:,} phao, cả hồ (không chia theo bè)'.replace(',', '.')),
    ('Bộ bulông nhựa nối phao', 'bộ', lambda j: f'=ROUNDUP({dlref("FT", j)}*{G["K_FB"]},0)', 'hh', 'ƯL', 'Tổng phao bè × hệ số K_FB ở GiaDinh'),
    ('B. DẦM C52 (đặt trên rãnh phao)',),
    ('Dầm C52 — số thanh liên tục', 'thanh', lambda j: '=' + dlref('C52N', j), 'same', 'CAD', 'Đếm trên bản vẽ CAD khung kết cấu (đối chiếu bảng tờ CT-KC-BE: DoiChieu_KhungThep)'),
    ('Dầm C52 — chiều dài', 'm', lambda j: '=' + dlref('C52L', j), 'same', 'CAD', 'Cộng từng thanh trên bản vẽ CAD'),
    ('Dầm C52 — số cây cấp hàng', 'cây', lambda j: '=' + barsum('C52_TungThanh', 'D', RAFTS[j]), 'same', 'CAD', 'Từng thanh: ROUNDUP(L / L_C52) (sheet C52_TungThanh)'),
    ('Bộ nối dầm C52 (Splice for Beam-C)', 'bộ', lambda j: '=' + barsum('C52_TungThanh', 'E', RAFTS[j]), 'hh', 'CAD', 'Số cây − số thanh. Nối đối đầu; mỗi bộ: 1 bản nối + 2 bulông M10×70 + 2 chốt M10×60 + đai ốc, long đen (tài liệu phụ kiện Mibet)'),
    ('Dầm C52 — khối lượng', 'kg', lambda j: f'={dlref("C52L", j)}*{G["M_C52"]}', 'same', 'ƯL', 'Theo tiết diện GIẢ ĐỊNH'),
    ('Bulông M8×35 bắt dầm C52 vào phao', 'bộ', lambda j: f'=ROUNDUP({dlref("C52L", j)}*{G["K_XF"]},0)', 'hh', 'ƯL', '1 bộ tại mỗi chỗ dầm C52 vượt qua một hàng phao (K_XF)'),
    ('C. HỘP 40x60 KẾT CẤU (nằm ngang trên dầm C52)',),
    ('Hộp 40x60 — số thanh liên tục', 'thanh', lambda j: '=' + dlref('HN', j), 'same', 'CAD', 'Đếm trên bản vẽ CAD khung kết cấu'),
    ('Hộp 40x60 — chiều dài', 'm', lambda j: '=' + dlref('HL', j), 'same', 'CAD', 'Cộng từng thanh trên bản vẽ CAD'),
    ('Hộp 40x60 — số cây cấp hàng', 'cây', lambda j: '=' + barsum('Hop40x60_TungThanh', 'D', RAFTS[j]), 'same', 'CAD', 'Từng thanh: ROUNDUP(L / L_H) (sheet Hop40x60_TungThanh)'),
    ('Bộ nối hộp 40x60 (Splice for Square Steel)', 'bộ', lambda j: '=' + barsum('Hop40x60_TungThanh', 'E', RAFTS[j]), 'hh', 'CAD', 'Nối đối đầu bằng bộ nối của nhà sản xuất (datasheet G5G)'),
    ('Hộp 40x60 — khối lượng', 'kg', lambda j: f'={dlref("HL", j)}*{G["M_H"]}', 'same', 'ƯL', 'Theo tiết diện GIẢ ĐỊNH'),
    ('Bulông M10×120 bắt hộp 40x60 vào dầm C52', 'bộ', lambda j: f'=ROUNDUP({dlref("HL", j)}/{G["B_C52"]},0)', 'hh', 'ƯL', '1 bộ tại mỗi giao điểm hộp – dầm C52 (HD lắp G5 bước 8)'),
    ('D. HỘP 40x60 GIA CƯỜNG TẠI VỊ TRÍ NEO (vẽ trên bản vẽ khung kết cấu, lớp GC-*; sức chịu móc "几" GIẢ ĐỊNH)',),
    ('Hộp 40x60 gia cường — số thanh', 'thanh', lambda j: '=' + gcsum('K', RAFTS[j]), 'same', 'CAD', 'Sheet GiaCuong_TaiNeo (từng cụm, mã cụm như trên bản vẽ khung có gia cường)'),
    ('Hộp 40x60 gia cường — chiều dài', 'm', lambda j: '=' + gcsum('M', RAFTS[j]), 'same', 'CAD', 'Đo trên bản vẽ khung có gia cường (thanh dài 15 m, ngắn hơn ở cụm sát nhau)'),
    ('Hộp 40x60 gia cường — số cây cấp hàng', 'cây', lambda j: '=' + gcsum('N', RAFTS[j]), 'same', 'CAD', 'Từng thanh: ROUNDUP(L / L_H)'),
    ('Bộ nối hộp 40x60 (gia cường)', 'bộ', lambda j: '=' + gcsum('O', RAFTS[j]), 'hh', 'ƯL', ''),
    ('Móc chữ "几" kẹp hộp vào dầm C52', 'cái', lambda j: '=' + gcsum('P', RAFTS[j]), 'hh', 'CAD', 'HD lắp G5 bước 8: hộp dùng làm điểm neo kẹp bằng móc "几" tại mọi giao điểm với C52 — đếm giao điểm trên bản vẽ'),
    ('Bộ đai ốc thoi M10×30 (cho móc "几")', 'bộ', lambda j: '=' + gcsum('Q', RAFTS[j]), 'hh', 'ƯL', '2 bộ mỗi móc'),
    ('Bulông M10×120 (hộp gia cường)', 'bộ', lambda j: '=' + gcsum('R', RAFTS[j]), 'hh', 'ƯL', '1 bộ mỗi móc'),
    ('Hộp 40x60 gia cường — khối lượng', 'kg', lambda j: '=' + gcsum('S', RAFTS[j]), 'same', 'ƯL', 'Theo tiết diện GIẢ ĐỊNH'),
    ('E. GÁ ĐỠ VÀ KẸP TẤM PIN',),
    ('Gá trước (Front Support G5)', 'cái', lambda j: f'={dlref("PV", j)}*{G["K_FS"]}', 'hh', 'ƯL', ''),
    ('Gá sau (Back Support G5)', 'cái', lambda j: f'={dlref("PV", j)}*{G["K_BS"]}', 'hh', 'ƯL', ''),
    ('Bộ kẹp biên tấm pin (End Clamp Kit, bulông M8×30)', 'bộ', lambda j: f'={dlref("PV", j)}*{G["K_EC"]}', 'hh', 'ƯL', ''),
    ('Bộ kẹp tiếp địa tấm pin', 'bộ', lambda j: f'={dlref("PV", j)}*{G["K_GC"]}', 'hh', 'ƯL', 'Cần nối đẳng thế khung tấm pin; hỏi Mibet kẹp của MRac G5 có răng tiếp địa không và định mức bộ/tấm trước khi đặt'),
    ('Bộ đai ốc thoi M8×25 + long đen lớn (gá vào C52)', 'bộ', lambda j: f'={dlref("PV", j)}*{G["K_RN8"]}', 'hh', 'ƯL', ''),
    ('F. ĐIỂM NEO TRÊN BÈ',),
    ('Ma ní chữ U (U-Shaped Shackle) làm bích neo trên bè', 'cái', lambda j: f'=2*{dlref("LINE", j)}', 'hh', 'NEO', f'Cáp 2 nhánh chữ Y: 2 bích mỗi dây, bắt vào hộp 40x60 gia cường; mắt nhánh cáp móc thẳng vào ma ní này. Số dây theo mặt bằng neo của web ({len(lines)} dây)'),
]
row = 4; stt = 0; item_rows = {}
NCOL = 8 + NR
for it in ITEMS:
    if len(it) == 1:
        th.cell(row, 2, it[0]).font = FB
        for c in range(1, NCOL + 1): th.cell(row, c).fill = SEC; th.cell(row, c).border = BORDER
        row += 1; continue
    name, unit, fn, total, src, note = it
    stt += 1; r = row; item_rows[name] = r
    th.cell(r, 1, stt); th.cell(r, 2, name); th.cell(r, 4, unit); th[f'{T_SRC}{r}'] = src; th[f'{T_NOTE}{r}'] = note
    if total == 'cable':
        th[f'{T_TOT}{r}'] = f'=DuLieuBanVe!{D_TOT}{D["FC"]}'
    else:
        for j in range(NR): th[f'{TC[j]}{r}'] = fn(j)
        th[f'{T_TOT}{r}'] = f'=SUM({TC[0]}{r}:{TC[-1]}{r})'
    th[f'{T_HH}{r}'] = f'=ROUNDUP({T_TOT}{r}*(1+{G["HH"]}/100),0)' if total == 'hh' else f'={T_TOT}{r}'
    for col in TC + [T_TOT, T_HH]: th[f'{col}{r}'].number_format = '#,##0'
    row += 1
grid(th, 3, row - 1, NCOL)
for r in range(4, row):
    for col in 'AD': th[f'{col}{r}'].alignment = CENTER
for j, w in enumerate([6, 46, 23, 8] + [10] * NR + [12, 13, 8, 80]): th.column_dimensions[COL(j + 1)].width = w
th.freeze_panes = 'E4'

tpl = openpyxl.load_workbook(TEMPLATE)['TongHop']
pics = {}
for im in tpl._images:
    name = tpl.cell(im.anchor._from.row + 1, 2).value
    a = im.anchor
    w, h = ((a.to.colOff - a._from.colOff) / 9525, (a.to.rowOff - a._from.rowOff) / 9525) if hasattr(a, 'to') and a.to.col == a._from.col else (im.width, im.height)
    pics[name] = (im._data(), w, h)
for name, r in item_rows.items():
    src = pics.get(name)
    if src:
        img = XLImage(io.BytesIO(src[0])); img.width, img.height = src[1], src[2]
        th.add_image(img, f'C{r}'); th.row_dimensions[r].height = 78
    else:
        th.cell(r, 3, '(không có hình trong tài liệu)').alignment = CENTER
        th.row_dimensions[r].height = 30

# ---- DoiChieu_KhungThep
dk = wb.create_sheet('DoiChieu_KhungThep')
dk['A1'] = 'ĐỐI CHIẾU BẢNG "THỐNG KÊ KHUNG THÉP" TRÊN TỜ CT-KC-BE-01…06 VỚI KHUNG VẼ TRONG CHÍNH TỜ ĐÓ (đếm và đo từng thanh)'; dk['A1'].font = FB
for j, h in enumerate(['Bè', 'Hạng mục', 'Số thanh — bảng', 'Số thanh — CAD', 'Lệch', 'Chiều dài — bảng (m)', 'Chiều dài — CAD (m)', 'Lệch (m)', 'Kết luận']):
    c = dk.cell(3, j + 1, h); c.font = FB; c.fill = HEAD; c.alignment = CENTER
r = 4
for j, n in enumerate(RAFTS):
    for item, nb, lb, nc, lc in (('Dầm C52', 'C52N_BV', 'C52L_BV', 'C52N', 'C52L'), ('Hộp 40x60', 'HN_BV', 'HL_BV', 'HN', 'HL')):
        col = RC[j]
        dk.cell(r, 1, n); dk.cell(r, 2, item)
        dk[f'C{r}'] = f'=DuLieuBanVe!{col}{D[nb]}'; dk[f'D{r}'] = f'=DuLieuBanVe!{col}{D[nc]}'; dk[f'E{r}'] = f'=D{r}-C{r}'
        dk[f'F{r}'] = f'=DuLieuBanVe!{col}{D[lb]}'; dk[f'G{r}'] = f'=DuLieuBanVe!{col}{D[lc]}'; dk[f'H{r}'] = f'=G{r}-F{r}'
        dk[f'I{r}'] = f'=IF(AND(E{r}=0,ABS(H{r})<=0.05),"KHỚP",IF(AND(E{r}=0,ABS(H{r})<=0.01*F{r}),"KHỚP (lệch nhỏ)","KHÔNG KHỚP bảng tờ — lấy theo bản vẽ CAD"))'
        for c_ in 'FGH': dk[f'{c_}{r}'].number_format = '#,##0.00'
        r += 1
grid(dk, 3, r - 1, 9)
for col, w in zip('ABCDEFGHI', [8, 12, 11, 11, 8, 14, 14, 10, 40]): dk.column_dimensions[col].width = w
DK_LAST = r - 1

# ---- GhiChu
gh = wb.create_sheet('GhiChu')
gh['A1'] = 'GHI CHÚ, NGUỒN SỐ LIỆU VÀ CÁC ĐIỂM CẦN XÁC NHẬN'; gh['A1'].font = Font(name='Calibri', size=13, bold=True)
NOTES = [
    'Nguồn: "BO TRI BE PIN HO BAN MONG.pdf" (Revit, Chủ đầu tư 30/10/2026: CT-BP-01…06 bố trí từng bè có bảng tấm pin, phao, inverter; MB-B01-01 mặt bằng tổng có bảng phao đỡ cáp) và 6 tờ khung kết cấu "CHUGIAP-HOBANMONG ver 2-Sheet - CT-KC-BE-01…06 - BỐ TRÍ KHUNG KẾT CẤU BÈ PIN 0n.dxf"; HD lắp Mibet MRac G5, tài liệu phụ kiện Mibet, datasheet G5G. Script: scripts/buildHuoiVanhRaftReinforcement.py --lake=ban-mong, scripts/buildBanMongRaftBOQ.py.',
    'Số tấm pin, phao, inverter lấy NGUYÊN bảng Revit trên tờ CT-BP (script kiểm tổng từng bảng). Dầm C52, hộp 40x60 kết cấu ĐO từng thanh trên bản vẽ CAD khung có gia cường (sheet C52_TungThanh, Hop40x60_TungThanh); đối chiếu bảng trên tờ khung ở sheet DoiChieu_KhungThep.',
    f'Số tấm pin {NPV:,} tấm (BÈ 1: 2.912; BÈ 2: 4.200; BÈ 3: 700; BÈ 4: 2.100; BÈ 5: 2.800; BÈ 6: 4.200). Từ 30/10/2026 lực gió, lực dây, cáp, cọc, đế trên web và trong thuyết minh, bảng tính, bảng vật tư neo đều tính theo số tấm này (trước đó tạm 16.900 chia theo diện tích).'.replace(',', '.'),
    'Hộp 40x60 gia cường tại vị trí neo: vẽ thêm vào bản vẽ khung (lớp GC-HOP 40x60 GIA CUONG, GC-BICH NEO, GC-DAY NEO, GC-GHI CHU; file "... BỐ TRÍ KHUNG KẾT CẤU BÈ PIN 0n - GIA CƯỜNG TẠI VỊ TRÍ NEO.dwg"). Cáp 2 nhánh chữ Y, 2 bích mỗi dây cách nhau 3,5 m, mỗi bích chịu K_Y × T; T là lực dây của bè trong bảng vật tư neo (engine, gió 20 m/s, bù mực nước liên tục + khóa bão). Với P_KEP = 7 kN GIẢ ĐỊNH, cụm chịu lực lớn cần 2–6 thanh trên các hàng phao kế nhau — cần số liệu sức chịu móc "几" của Mibet; chi tiết liên kết bích – hộp do nhà cung cấp bè lập.',
    'Mặt bằng neo (số dây, vị trí bích) là của web (ưu tiên bờ, tầm dây bờ 80 m, 243 dây), tính trên mặt bằng bè CAD 30/10/2026. Bản vẽ khung là mặt bằng đó xoay quanh gốc tọa độ (BÈ 1 105°, BÈ 2 15°, BÈ 3 94°, BÈ 4 84,9°, BÈ 5 60,3°, BÈ 6 35°), khớp đường bao lệch trung bình 0,3–0,8 m; khung rộng hơn đường bao 1,3–2 m ở một cạnh nên vị trí bích có thể lệch ~1 m dọc mép (thanh 15 m vẫn phủ).',
    'Khối lượng thép theo tiết diện GIẢ ĐỊNH (C52 ≈ 2,36 kg/m; hộp 40×60×2,0 ≈ 3,01 kg/m) vì chưa có catalogue tiết diện của nhà sản xuất.',
    'Gá đỡ, kẹp tấm pin, bulông nối phao, bulông M8 / M10 của khung là ƯỚC LƯỢNG theo hệ số mỗi tấm / mỗi phao / mỗi mét (sheet GiaDinh) — phải thay bằng BOM của nhà sản xuất khi có. Kẹp tiếp địa: hỏi Mibet trước khi đặt.',
    'Chưa gồm: cáp neo, phụ kiện đầu cáp, thiết bị bù mực nước, cọc khoan nhồi, đế neo (xem bảng thống kê vật tư neo Bản Mòng), cáp điện DC/AC, máng cáp, tủ điện, tấm lối đi, vận chuyển.',
    'Hình minh họa ở cột C chép từ bảng vật tư bè pin Huổi Vanh trước đây (cắt từ tài liệu Mibet): chỉ để nhận dạng loại vật tư, không phải bản vẽ chế tạo.',
]
for i, t in enumerate(NOTES):
    gh.cell(3 + i, 1, i + 1); gh.cell(3 + i, 2, t).alignment = WRAP
grid(gh, 3, 2 + len(NOTES), 2)
gh.column_dimensions['A'].width = 5; gh.column_dimensions['B'].width = 150

wb._sheets = [wb['TongHop'], wb['DuLieuBanVe'], wb['GiaCuong_TaiNeo'], wb['C52_TungThanh'], wb['Hop40x60_TungThanh'], wb['DoiChieu_KhungThep'], wb['DoiChieu_MatBang'], wb['GiaDinh'], wb['GhiChu']]
wb.active = 0
wb.save(OUT)
add_apply_flags(OUT)

# ------------------------------------------------------------------ drawing GC-01
def clip_line(pts, p, d):
    ts = []
    for (x1, y1), (x2, y2) in zip(pts, pts[1:] + pts[:1]):
        ex, ey = x2 - x1, y2 - y1
        den = d[0] * ey - d[1] * ex
        if abs(den) < 1e-12: continue
        t = ((x1 - p[0]) * ey - (y1 - p[1]) * ex) / den
        u = ((x1 - p[0]) * d[1] - (y1 - p[1]) * d[0]) / den
        if 0 <= u < 1: ts.append(t)
    ts.sort()
    return [(ts[i], ts[i + 1]) for i in range(0, len(ts) - 1, 2)]


dxf = ezdxf.new('R2010', setup=True)
dxf.header['$INSUNITS'] = 6
dxf.styles.add('VN', font='arial.ttf')
for lname, color, lw in (('BE_PIN', 7, 50), ('KHUNG_C52', 8, 0), ('DAY_NEO', 5, 0), ('BICH_NEO', 1, 25), ('GIA_CUONG_40x60', 1, 70), ('GHI_CHU', 7, 0)):
    dxf.layers.add(lname, color=color, lineweight=lw)
msp = dxf.modelspace()
vn = lambda v, d=0: f'{v:,.{d}f}'.replace(',', ' ').replace('.', ',').replace(' ', '.')
txt = lambda s, at, h, layer='GHI_CHU', rot=0: msp.add_text(s, height=h, rotation=rot, dxfattribs={'layer': layer, 'style': 'VN'}).set_placement(at)
for name in RAFTS:
    pts = [(p['x'], p['y']) for p in polys[name]['points']]
    msp.add_lwpolyline(pts, close=True, dxfattribs={'layer': 'BE_PIN'})
    a = math.radians(RJ[name]['fit']['rot_deg'])
    fb = (1.0, 0.0) if RJ[name]['boxAlongX'] else (0.0, 1.0)
    bdir = (math.cos(a) * fb[0] + math.sin(a) * fb[1], -math.sin(a) * fb[0] + math.cos(a) * fb[1]); cdir = (-bdir[1], bdir[0])
    proj = [x * bdir[0] + y * bdir[1] for x, y in pts]
    s = min(proj) + B_C52 / 2
    while s < max(proj):
        p = (bdir[0] * s, bdir[1] * s)
        for t0, t1 in clip_line(pts, p, cdir):
            msp.add_line((p[0] + cdir[0] * t0, p[1] + cdir[1] * t0), (p[0] + cdir[0] * t1, p[1] + cdir[1] * t1), dxfattribs={'layer': 'KHUNG_C52'})
        s += 5 * B_C52
    c = polys[name]['centroid']
    txt(f'{name} — T = {vn(T[name], 1)} kN, {lineCount[name]} dây', (c['x'] - 30, c['y']), 4.0, 'BE_PIN')
for ln in lines:
    msp.add_line((ln['xRaft'], ln['yRaft']), (ln['xAnchor'], ln['yAnchor']), dxfattribs={'layer': 'DAY_NEO'})
for name, cid, kind, codes, c in clusters:
    for b in c['bars']:
        msp.add_lwpolyline(b['world'], dxfattribs={'layer': 'GIA_CUONG_40x60', 'const_width': 0.3})
    for (bx, by) in c['cleatsWorld']:
        msp.add_circle((bx, by), 0.6, dxfattribs={'layer': 'BICH_NEO'})
    (x0, y0), (x1, y1) = c['bars'][-1]['world']
    ang = math.degrees(math.atan2(y1 - y0, x1 - x0))
    if ang > 90: ang -= 180
    if ang < -90: ang += 180
    txt(f"{cid}: {c['n']} thanh", ((x0 + x1) / 2, (y0 + y1) / 2), 1.2, 'GIA_CUONG_40x60', ang)
xs = [p['x'] for n in RAFTS for p in polys[n]['points']]; ys = [p['y'] for n in RAFTS for p in polys[n]['points']]
X0, Y0 = max(xs) + 60, max(ys)
txt('SƠ ĐỒ TỔNG HỘP 40x60 GIA CƯỜNG TẠI CÁC VỊ TRÍ NEO — HỒ BẢN MÒNG, 6 BÈ (GC-01)', (X0, Y0), 6.0)
txt(f'Thanh lấy từ bản vẽ khung có gia cường "CT-KC-BE-0n ... - GIA CƯỜNG TẠI VỊ TRÍ NEO.dwg" (chi tiết từng thanh, từng móc); mặt bằng neo: web ({len(lines)} dây, cáp chữ Y). Đơn vị: m.', (X0, Y0 - 10), 3.0)
y = Y0 - 22
for n in RAFTS:
    cl = [c for c in clusters if c[0] == n]
    nb = sum(c[4]['n'] for c in cl); L = sum(b['L_m'] for c in cl for b in c[4]['bars']); nk = sum(b['nC52'] for c in cl for b in c[4]['bars'])
    txt(f'{n}: {len(cl)} cụm, {nb} thanh, {vn(L)} m hộp 40x60 gia cường, {vn(nk)} móc "几"; T = {vn(T[n], 1)} kN', (X0, y), 3.0); y -= 6
tot = sum(c[4]['n'] for c in clusters)
Lt = sum(b['L_m'] for c in clusters for b in c[4]['bars'])
txt(f'TỔNG: {len(clusters)} cụm, {tot} thanh, {vn(Lt)} m', (X0, y), 3.5); y -= 10
for t in ['GHI CHÚ:',
          '1. Mỗi dây neo có 2 bích (cáp 2 nhánh chữ Y) cách nhau 3,5 m, vòng tròn đỏ. Mã cụm và số thanh: sheet GiaCuong_TaiNeo của bảng vật tư bè pin.',
          '2. Thanh nằm trên hàng phao, theo phương hộp 40x60 kết cấu. Bích trên cạnh song song hộp: thanh dọc cạnh trên hàng phao thứ 2 tính từ mép;',
          '   bích trên cạnh vuông góc: mỗi bích một thanh từ dầm C52 ngoài cùng chạy vào trong.',
          '3. Hộp kẹp vào MỌI dầm C52 nó cắt qua bằng móc "几" (2 đai ốc thoi M10×30 + 1 bulông M10×120).',
          '4. Số thanh = MAX(1; số bích × 0,7 T × 1,2 / (số dầm C52 thanh cắt qua × 7 kN)); sức chịu móc 7 kN là GIẢ ĐỊNH, chờ số liệu Mibet.',
          '5. Nét xám mảnh: phương dầm C52 (vẽ 1/5 số dầm). Nét xanh: dây neo. Chi tiết liên kết bích – hộp do nhà cung cấp bè lập và xác nhận.',
          '6. Vị trí bích lấy theo mặt bằng neo; khung trên bản vẽ rộng hơn đường bao bè 1,3–2 m ở một cạnh nên bích có thể lệch ~1 m dọc mép.']:
    txt(t, (X0, y), 2.6); y -= 5
dxf.saveas(OUT_DXF)

print(json.dumps({'out': OUT, 'dxf': OUT_DXF, 'clusters': len(clusters), 'bars': tot, 'length_m': round(Lt, 1), 'panels': NPV, 'inverters': NINV,
                  'sheets': SHEETS, 'cableFloats': CABLE_FLOATS,
                  'cad': {n: [len(C52_CAD[n]), round(sum(C52_CAD[n]), 2), len(BOX_CAD[n]), round(sum(BOX_CAD[n]), 2)] for n in RAFTS}}, ensure_ascii=False))
