"""
Trình bày bảng neo do scripts/buildMooringOrderBOQ.ts xuất ra giống bảng của
Huổi Vanh: thêm cột C "HÌNH MINH HỌA" ở sheet tổng hợp (hình chép từ file mẫu
theo tên hạng mục, bỏ cỡ cáp / WLL / đường kính khi so tên), kẻ khung, tô dòng
tiêu đề và dòng mục, chữ xuống dòng.

  python scripts/styleMooringBOQ.py <bảng neo .xlsx> <mẫu Huổi Vanh .xlsx> <out.xlsx>
"""
import io, re, sys
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from xlsx_apply_flags import add_apply_flags

SRC, TEMPLATE, OUT = sys.argv[1:4]
thin = Side(style='thin'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='F2F2F2'); SEC = PatternFill('solid', fgColor='D9D9D9')
F = Font(name='Calibri', size=12); FB = Font(name='Calibri', size=12, bold=True)
WRAP = Alignment(wrap_text=True, vertical='center'); CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIAS = {  # Bản Mòng rows with no Huổi Vanh twin of the same name -> the Huổi Vanh row with the same part
    'tai neo đơn mặt đế': 'tai neo đôi mặt đế',
    'bản mã đầu cọc mạ kẽm (cọc đơn và đài cọc đôi)': 'bản nắp thép đầu cọc mạ kẽm',
    'tai neo đầu cọc / đài (padeye)': 'tai neo đầu cọc (padeye đơn)',
    'thép neo bản mã vào bê tông': 'râu thép ngàm bản nắp',
    'tai neo mặt đế (bản mã + tai + thép neo 1 bộ mỗi dây)': 'tai neo đôi mặt đế',
    'thép tấm bản mã + tai neo (đầu cọc đài đế)': 'thép tấm bản nắp 400×400×20',
    'thép neo bản mã': 'thép râu ngàm bản nắp',
}

def key(name):
    s = str(name).lower()
    s = re.sub(r'\(wll [^)]*\)|wll [\d.,]+ t|pes-\d+|cb\d+-v|[φø]\d+(×\d+)?', '', s)
    s = re.sub(r'[\s,]+', ' ', s).strip()
    return ALIAS.get(s, s)

tpl = openpyxl.load_workbook(TEMPLATE).worksheets[0]
pics = {}
for im in tpl._images:
    a = im.anchor  # shown size = the template's two-cell anchor (the picture inside column C)
    w, h = ((a.to.colOff - a._from.colOff) / 9525, (a.to.rowOff - a._from.rowOff) / 9525) if hasattr(a, 'to') and a.to.col == a._from.col else (im.width, im.height)
    pics.setdefault(key(tpl.cell(a._from.row + 1, 2).value), (im._data(), w, h))

wb = openpyxl.load_workbook(SRC)
ws = wb.worksheets[0]
ws.insert_cols(3)
HEADER_ROW = next(r for r in range(1, 10) if ws.cell(r, 1).value == 'STT')
ws.cell(HEADER_ROW, 3, 'HÌNH MINH HỌA')
last = ws.max_row
for r in range(1, HEADER_ROW - 1):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    ws.cell(r, 1).alignment = WRAP
ws['A1'].font = Font(name='Calibri', size=13, bold=True)
ws.row_dimensions[2].height = 34
for r in range(HEADER_ROW, last + 1):
    a = ws.cell(r, 1).value
    is_sec = isinstance(a, str) and a != 'STT'
    for c in range(1, 10):
        cell = ws.cell(r, c); cell.border = BORDER
        cell.font = FB if (r == HEADER_ROW or is_sec) else F
        cell.alignment = CENTER if (r == HEADER_ROW or c in (1, 3, 5, 7)) else WRAP
        if r == HEADER_ROW: cell.fill = HEAD
        elif is_sec: cell.fill = SEC
    if r == HEADER_ROW or is_sec:
        ws.row_dimensions[r].height = 31; continue
    name = ws.cell(r, 2).value
    pic = pics.get(key(name))
    if pic:
        img = XLImage(io.BytesIO(pic[0])); img.width, img.height = pic[1], pic[2]
        ws.add_image(img, f'C{r}'); ws.row_dimensions[r].height = 92
    else:
        ws.cell(r, 3, '(chưa có hình)'); ws.row_dimensions[r].height = 46
    for c in (6, 8): ws.cell(r, c).number_format = '#,##0.0' if ws.cell(r, 5).value in ('md', 'm³') else '#,##0'
for col, w in zip('ABCDEFGHI', [6.8, 44.8, 24.6, 52.8, 7.8, 14.8, 12.8, 14.8, 70.8]):
    ws.column_dimensions[col].width = w

for other in wb.worksheets[1:]:
    other['A1'].font = Font(name='Calibri', size=13, bold=True)
    hdr = next((r for r in range(2, 6) if sum(1 for c in other[r] if c.value not in (None, '')) >= 3), None)
    if hdr is None: continue
    for row in other.iter_rows(min_row=hdr, max_row=other.max_row):
        if all(c.value in (None, '') for c in row): continue
        for c in row:
            c.border = BORDER; c.font = FB if c.row == hdr else F
            c.alignment = CENTER if c.row == hdr else Alignment(vertical='center')
            if c.row == hdr: c.fill = HEAD
    other.row_dimensions[hdr].height = 48
    other.freeze_panes = other.cell(hdr + 1, 1)
wb.save(OUT)
add_apply_flags(OUT)
print(OUT, 'images', len(ws._images), 'rows', last)
